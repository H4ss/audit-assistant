"""Construction déterministe du classeur cible de démo.

Le classeur est généré plutôt que commité (aucun binaire métier dans Git).
Il contient volontairement :
  * un onglet `Fortify` aux en-têtes du contrat (section 21.1), sans lignes ;
  * un onglet `ToolB` aux en-têtes hétérogènes, déjà rempli par l'inventaire
    du concurrent, avec une colonne libre de l'équipe, une formule par ligne
    et une décision humaine préexistante à préserver ;
  * un onglet `Synthèse` contenant des formules et une mise en forme ;
  * pas d'onglet pour ToolC (nouvel outil).
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from paladin.contracts import DEFAULT_COLUMNS, criticality_in_header, found_in_header

TOOLB_HEADERS = [
    "Finding ID",
    "Application",
    "Rule",
    "File path",
    "Line",
    "Severity",
    "CWE",
    "Commentaires",
    "Notes équipe",
    "analysis result",
    "Analysis result comment",
    found_in_header("Fortify"),
    criticality_in_header("Fortify"),
    "Suivi",
]

# Inventaire ToolB (ligne Excel). TB-0009 n'existe que dans le MD ; TB-0007
# n'existe que dans l'Excel ; TB-0003 diverge (High ici, Critical dans le MD).
TOOLB_ROWS = [
    ("TB-0001", "ShopApp", "TB.SQLI.001", "shop-api/app/orders.py", 12, "High", "CWE-89", None, None),
    ("TB-0002", "ShopApp", "TB.SQLI.001", "shop-api/app/products.py", 11, "High", "CWE-89", None, None),
    (
        "TB-0003",
        "ShopApp",
        "TB.XSS.004",
        "shop-api/app/search.py",
        7,
        "High",
        "CWE-79",
        None,
        "voir avec l'équipe front",
    ),
    ("TB-0004", "ShopApp", "TB.XSS.004", "shop-api/app/search.py", 12, "Medium", "CWE-79", None, None),
    ("TB-0005", "ShopApp", "TB.LOG.002", "shop-api/app/logging_utils.py", 7, "Medium", "CWE-117", None, None),
    ("TB-0006", "ShopApp", "TB.RAND.001", "billing-lib/billing/tokens.py", 8, "High", "CWE-338", None, None),
    (
        "TB-0007",
        "ShopApp",
        "TB.REDIR.002",
        "shop-api/app/orders.py",
        18,
        "Low",
        "CWE-601",
        "Import ToolB 2026-09",
        None,
    ),
    ("TB-0008", "ShopApp", "TB.CRYPTO.007", "billing-lib/billing/hashing.py", 7, "Low", "CWE-328", None, None),
]

# Décision saisie à la main avant Paladin : doit être préservée.
TOOLB_HUMAN_DECISION = {"TB-0008": ("Not an issue", "cache key only, no security use")}


def fortify_headers(other_tools: list[str]) -> list[str]:
    headers = [c.header for c in DEFAULT_COLUMNS]
    for tool in other_tools:
        headers += [found_in_header(tool), criticality_in_header(tool)]
    return headers


def build_demo_workbook(path: Path) -> Path:
    wb = Workbook()
    summary = wb.active
    summary.title = "Synthèse"
    summary["A1"] = "Synthèse de l'audit (démo)"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A3"] = "Fortify — True Positive"
    summary["B3"] = '=COUNTIF(Fortify!N:N,"True Positive")'
    summary["A4"] = "Fortify — Not an issue"
    summary["B4"] = '=COUNTIF(Fortify!N:N,"Not an issue")'
    summary["A5"] = "ToolB — True Positive"
    summary["B5"] = '=COUNTIF(ToolB!J:J,"True Positive")'
    summary["A6"] = "ToolB — Not an issue"
    summary["B6"] = '=COUNTIF(ToolB!J:J,"Not an issue")'
    summary["A8"] = "Rédigé par l'équipe AppSec — ne pas modifier"
    summary["A8"].fill = PatternFill("solid", fgColor="FFF2CC")
    summary.column_dimensions["A"].width = 36

    fortify = wb.create_sheet("Fortify")
    fortify.append(fortify_headers(["ToolB"]))
    for cell in fortify[1]:
        cell.font = Font(bold=True)
    fortify.freeze_panes = "A2"

    toolb = wb.create_sheet("ToolB")
    toolb.append(TOOLB_HEADERS)
    for cell in toolb[1]:
        cell.font = Font(bold=True)
    for i, (fid, app, rule, file_, line, sev, cwe, comments, notes) in enumerate(TOOLB_ROWS, start=2):
        result, comment = TOOLB_HUMAN_DECISION.get(fid, (None, None))
        toolb.append([fid, app, rule, file_, line, sev, cwe, comments, notes, result, comment, None, None])
        toolb.cell(row=i, column=14, value=f'=IF(J{i}="","à faire","fait")')
    toolb.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
