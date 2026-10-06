"""Lecture générique de sources tabulaires : Excel (fichier ou onglet) et CSV."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from paladin.contracts import Completeness
from paladin.importers.records import RawRecord, SourceRead, UnrecognizedPart

HEADER_SCAN_ROWS = 10


class TabularError(ValueError):
    pass


def _is_header_row(values: list[Any]) -> bool:
    texts = [v for v in values if isinstance(v, str) and v.strip()]
    return len(texts) >= 2 and len(texts) >= 0.6 * len([v for v in values if v not in (None, "")])


def _rows_to_read(rows: list[list[Any]], locator_prefix: str, kind: str, first_row_number: int = 1) -> SourceRead:
    header_idx = next((i for i, r in enumerate(rows[:HEADER_SCAN_ROWS]) if _is_header_row(r)), None)
    if header_idx is None:
        raise TabularError("Aucune ligne d'en-têtes reconnue dans les premières lignes.")
    raw_headers = rows[header_idx]
    headers: list[str] = []
    for h in raw_headers:
        name = str(h).strip() if h not in (None, "") else ""
        if name and name in headers:
            raise TabularError(f"En-tête en double : {name!r}. → Renommer la colonne dans la source.")
        headers.append(name)
    records: list[RawRecord] = []
    unrecognized: list[UnrecognizedPart] = []
    for offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 1):
        rownum = first_row_number + offset
        if all(v in (None, "") for v in row):
            continue
        fields = {h: (row[i] if i < len(row) else None) for i, h in enumerate(headers) if h}
        stray = [row[i] for i in range(len(row)) if (i >= len(headers) or not headers[i]) and row[i] not in (None, "")]
        if stray:
            unrecognized.append(
                UnrecognizedPart(f"{locator_prefix}R{rownum}", "valeurs hors colonnes nommées", str(stray)[:200])
            )
        records.append(RawRecord(fields=fields, locator=f"{locator_prefix}R{rownum}"))
    read = SourceRead(kind=kind, records=records, field_names=[h for h in headers if h], unrecognized=unrecognized)
    read.notes.append(f"En-têtes en ligne {first_row_number + header_idx}")
    read.scope["header_row"] = first_row_number + header_idx
    read.completeness = Completeness.PARTIAL if unrecognized else Completeness.COMPLETE
    return read


def list_sheets(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_excel(path: Path, sheet: str | None = None) -> SourceRead:
    """Lit un onglet (par défaut le premier) en valeurs brutes (formules non évaluées)."""
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet is None:
            sheet = wb.sheetnames[0]
        if sheet not in wb.sheetnames:
            raise TabularError(f"Onglet {sheet!r} absent. Onglets disponibles : {', '.join(wb.sheetnames)}")
        rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()
    read = _rows_to_read(rows, f"{path.name}!{sheet}!", "excel")
    read.scope["sheet"] = sheet
    return read


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise TabularError("Encodage du CSV non reconnu.")


def read_csv(path: Path) -> SourceRead:
    text = _decode(path.read_bytes())
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    return _rows_to_read(rows, f"{path.name}:", "csv")
