"""Création d'une campagne à partir d'un groupe SSC (`APP.*` = une entrée).

Chaque sous-application ayant une version `release` unique devient une version
collectée par la même source Fortify. Les sous-applications sans `release`
unique sont exclues et listées. Aucun classeur fourni : un classeur neuf est
généré à partir du contrat Excel.
"""

from __future__ import annotations

import sqlite3
from pathlib import Path
from typing import Any

from paladin import store
from paladin.config import Settings, ensure_campaign_dirs
from paladin.excel.template import new_workbook
from paladin.fortify.discovery import Group, campaign_slug


class SSCCampaignError(ValueError):
    pass


def descriptor_for_group(
    group: Group, campaign_id: str, name: str, target_workbook: str, repos: list[dict[str, Any]]
) -> dict[str, Any]:
    versions = [
        {"application_name": a.app_name, "application_id": a.app_id, "version_id": a.version_id} for a in group.ready
    ]
    return {
        "id": campaign_id,
        "name": name,
        "target_workbook": target_workbook,
        "repos": repos,
        "ssc_group": group.key,
        "excluded_subapps": [{"application_name": a.app_name, "reason": a.status} for a in group.blocked],
        "tools": [
            {
                "label": "Fortify",
                "kind": "fortify",
                "sheet_name": "Fortify",
                "sources": [{"role": "findings", "kind": "fortify_ssc"}],
                "fortify": {"version_name": "release", "versions": versions},
            }
        ],
    }


def create_campaign_from_group(
    settings: Settings,
    conn: sqlite3.Connection,
    group: Group,
    *,
    campaign_id: str | None = None,
    name: str | None = None,
    target_workbook: Path | None = None,
    repo_paths: list[Path] | None = None,
) -> str:
    if not group.ready:
        raise SSCCampaignError(f"Groupe {group.key} : aucune sous-application avec une version « release » unique.")
    cid = campaign_id or campaign_slug(group.key)
    if conn.execute("SELECT 1 FROM campaign WHERE id = ?", (cid,)).fetchone():
        raise SSCCampaignError(f"Une campagne « {cid} » existe déjà.")
    repos = []
    for p in repo_paths or []:
        p = p.expanduser().resolve()
        if not p.is_dir():
            raise SSCCampaignError(f"Dossier de dépôt introuvable : {p}")
        repos.append({"name": p.name, "path": str(p), "commit": None, "scanner_roots": []})
    root = ensure_campaign_dirs(settings, cid)
    if target_workbook is None:
        target = new_workbook(root / "inputs" / f"audit_{cid}.xlsx", {"Fortify": True})
    else:
        target = target_workbook.expanduser().resolve()
        if not target.is_file():
            raise SSCCampaignError(f"Classeur cible introuvable : {target}")
        from openpyxl import load_workbook

        wb = load_workbook(target, read_only=True)
        try:
            if "Fortify" not in wb.sheetnames:
                raise SSCCampaignError(
                    f"Le classeur {target.name} n'a pas d'onglet « Fortify ». Laisser vide pour en générer un neuf."
                )
        finally:
            wb.close()
    desc = descriptor_for_group(group, cid, name or f"{group.key} (Fortify release)", str(target), repos)
    with store.transaction(conn):
        store.create_campaign(conn, cid, desc["name"], desc)
        store.add_tool(conn, cid, "Fortify", "fortify", "Fortify")
        for r in repos:
            store.add_repo(conn, cid, r["name"], r["path"], r["commit"], r["scanner_roots"])
    return cid
