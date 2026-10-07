"""Classeur cible neuf, généré à partir du contrat Excel (section 21.1).

Utilisé quand aucun classeur n'est fourni (plug and play) : un onglet par outil,
en-têtes par défaut du contrat, paires comparatives pour les autres outils.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from paladin.contracts import DEFAULT_COLUMNS, FORTIFY_ONLY_KEYS, criticality_in_header, found_in_header


def headers_for(tool: str, all_tools: list[str], fortify: bool) -> list[str]:
    cols = [c.header for c in DEFAULT_COLUMNS if fortify or c.key not in FORTIFY_ONLY_KEYS]
    for other in all_tools:
        if other != tool:
            cols += [found_in_header(other), criticality_in_header(other)]
    return cols


def new_workbook(path: Path, tools: dict[str, bool]) -> Path:
    """`tools` : {libellé d'onglet: est_fortify}. N'écrase jamais un fichier existant."""
    if path.exists():
        raise FileExistsError(f"Le classeur existe déjà : {path}")
    wb = Workbook()
    wb.remove(wb.active)
    names = list(tools)
    for name, is_fortify in tools.items():
        ws = wb.create_sheet(name[:31])
        ws.append(headers_for(name, names, is_fortify))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
