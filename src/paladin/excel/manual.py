"""Import d'analyses manuelles depuis le classeur personnel de l'analyste.

Le classeur a son propre format : Paladin ne le modifie jamais (l'interface en
garde une copie). Les colonnes sont inférées à partir des en-têtes ET des
données : la colonne d'identifiant est celle dont les valeurs correspondent aux
findings importés, la colonne de verdict celle dont les valeurs ressemblent à
un verdict. Chaque ligne est rapprochée d'un finding par identifiant (Instance
ID), sinon par fichier + ligne (catégorie en départage). Rien n'est deviné en
silence : l'aperçu montre la lecture des colonnes et la traduction de chaque
valeur de verdict, toutes deux modifiables avant l'import.

Les lignes retenues deviennent des décisions « import » traçables, regroupées
sous un lot annulable. Une décision Paladin existante et différente est
conservée et listée en conflit.
"""

from __future__ import annotations

import re
import sqlite3
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from paladin import store
from paladin.contracts import Authority, DecisionAction, Verdict, verdict_to_excel
from paladin.importers.mapping import _SYNONYMS, apply_transform, norm_name, signature
from paladin.review import decisions
from paladin.review.decisions import excel_projection
from paladin.util import dumps, loads, new_id, utcnow

ROLES: dict[str, str] = {
    "key": "Identifiant (Instance ID)",
    "file": "Fichier",
    "line": "Ligne",
    "category": "Catégorie",
    "verdict": "Verdict",
    "comment": "Commentaire",
    "minutes": "Temps passé (minutes)",
}
IGNORE = "IGNORE"
VALUE_CHOICES = {
    Verdict.TRUE_POSITIVE.value: "True Positive",
    Verdict.NOT_AN_ISSUE.value: "Not an issue",
    IGNORE: "ignorer (pas une décision)",
}
_HEADERS: dict[str, set[str]] = {
    "key": {*_SYNONYMS["source_id"], "issue instance id", "fortify id", "instanceid"},
    "file": {*_SYNONYMS["full_filename"], "primary location", "primarylocation", "fichier source"},
    "line": set(_SYNONYMS["line_number"]),
    "category": {*_SYNONYMS["category"], "fortify category"},
    "verdict": {
        "analysis result",
        "verdict",
        "resultat",
        "resultat analyse",
        "resultat de l analyse",
        "statut",
        "status",
        "conclusion",
        "qualification",
        "analyse",
        "analysis",
        "tp fp",
        "decision",
        "audit",
    },
    "comment": {
        "analysis result comment",
        "comment",
        "comments",
        "commentaire",
        "commentaires",
        "justification",
        "remarque",
        "remarques",
        "notes",
        "note",
        "explication",
        "rationale",
        "analyse detaillee",
    },
    "minutes": {"temps", "temps passe", "temps min", "time", "time spent", "duree", "duration", "minutes", "min"},
}
# Traductions proposées (jamais appliquées sans être montrées) ; « ok », « oui », « non »… restent à préciser.
_VERDICT_GUESSES: dict[str, str] = {
    **dict.fromkeys(
        ("tp", "true positive", "vrai positif", "vp", "confirmed", "confirme", "exploitable", "vulnerable"),
        Verdict.TRUE_POSITIVE.value,
    ),
    **dict.fromkeys(
        (
            "fp",
            "false positive",
            "faux positif",
            "not an issue",
            "nai",
            "not issue",
            "no issue",
            "pas un probleme",
            "non exploitable",
        ),
        Verdict.NOT_AN_ISSUE.value,
    ),
}
_HEADER_SCAN_ROWS = 15
_PATHLIKE = re.compile(r"[/\\]|\.[A-Za-z0-9]{1,8}(:\d+)?$")


class ManualImportError(ValueError):
    pass


@dataclass
class ManualItem:
    row: int
    key: str | None
    file: str | None
    line: int | None
    raw_result: str | None
    comment: str | None
    minutes: float | None
    category: str | None = None
    finding_id: str | None = None
    source_id: str | None = None
    verdict: str | None = None
    matched_by: str | None = None  # identifiant | fichier + ligne
    status: str = (
        ""  # importable | identique | conflit | commentaire seul | non reconnu | sans finding | ambigu | doublon
    )
    detail: str = ""


@dataclass
class ManualPlan:
    path: Path
    sheet: str
    sheets: list[str]
    header_row: int
    headers: list[str]
    mapping: dict[str, int | None]
    reasons: dict[str, str]
    value_map: dict[str, str]
    value_counts: Counter
    signature: str
    profile_reused: bool = False
    items: list[ManualItem] = field(default_factory=list)

    def by_status(self, *statuses: str) -> list[ManualItem]:
        return [i for i in self.items if i.status in statuses]

    @property
    def importable(self) -> list[ManualItem]:
        return self.by_status("importable")

    @property
    def minutes_total(self) -> float | None:
        values = [i.minutes for i in self.items if i.minutes is not None and i.status in ("importable", "identique")]
        return sum(values) if values else None

    def column_label(self, role: str) -> str:
        i = self.mapping.get(role)
        return self.headers[i] if i is not None and i < len(self.headers) else "—"


def _text(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def _number(v: Any) -> float | None:
    if isinstance(v, int | float) and not isinstance(v, bool):
        return float(v)
    s = _text(v)
    if s is None:
        return None
    m = re.fullmatch(r"(\d+(?:[.,]\d+)?)\s*(min|mn|m)?", s.lower())
    return float(m.group(1).replace(",", ".")) if m else None


def guess_verdict(raw: Any) -> str | None:
    s = _text(raw)
    return _VERDICT_GUESSES.get(norm_name(s)) if s else None


def _norm_path(p: str | None) -> str | None:
    return p.replace("\\", "/").strip().lstrip("./").lower() if p else None


def _same_file(a: str, b: str) -> bool:
    return a == b or a.endswith("/" + b) or b.endswith("/" + a)


def _findings(conn: sqlite3.Connection, campaign_id: str, tool_label: str | None) -> list[dict[str, Any]]:
    sql = (
        "SELECT f.id, f.source_id, f.normalized_path, f.full_filename, f.line_number, f.category, f.revision,"
        " f.current_decision_id, t.label AS tool FROM finding f JOIN tool t ON t.id = f.tool_id WHERE f.campaign_id = ?"
    )
    params: list[Any] = [campaign_id]
    if tool_label:
        sql += " AND t.label = ?"
        params.append(tool_label)
    return [dict(r) for r in conn.execute(sql, params)]


def _read_sheet(path: Path, sheet: str | None) -> tuple[list[str], str, list[tuple]]:
    if not path.is_file():
        raise ManualImportError(f"Classeur introuvable : {path}")
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ManualImportError("Format attendu : classeur Excel .xlsx (enregistrer sous… .xlsx si besoin).")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # classeur corrompu, chiffré ou non Excel
        raise ManualImportError(f"Classeur illisible : {exc}") from None
    try:
        names = list(wb.sheetnames)
        if sheet and sheet not in names:
            raise ManualImportError(f"Onglet « {sheet} » absent ; onglets : {', '.join(names)}")
        data = {n: list(wb[n].iter_rows(values_only=True)) for n in ([sheet] if sheet else names)}
    finally:
        wb.close()
    if sheet:
        return names, sheet, data[sheet]
    return names, "", [data[n] for n in names]  # type: ignore[return-value]


def _header_row(rows: list[tuple]) -> int:
    known = set().union(*_HEADERS.values())
    best, best_score = 0, -1
    for i, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        cells = [norm_name(c) for c in row if _text(c)]
        score = sum(1 for c in cells if c in known)
        if score > best_score and cells:
            best, best_score = i, score
    return best


def _infer(
    headers: list[str], body: list[tuple], findings: list[dict[str, Any]]
) -> tuple[dict[str, int | None], dict[str, str]]:
    names = [norm_name(h) for h in headers]
    keys = {f["source_id"].lower() for f in findings}
    basenames = {
        (f["normalized_path"] or f["full_filename"] or "").replace("\\", "/").split("/")[-1].lower() for f in findings
    } - {""}
    columns = [[row[i] if i < len(row) else None for row in body] for i in range(len(headers))]
    mapping: dict[str, int | None] = dict.fromkeys(ROLES)
    reasons: dict[str, str] = {}
    taken: set[int] = set()

    def pick(role: str, i: int, why: str) -> None:
        mapping[role] = i
        reasons[role] = why
        taken.add(i)

    hits = [sum(1 for v in col if (_text(v) or "").lower() in keys) for col in columns]
    if hits and max(hits) > 0:
        i = max(range(len(hits)), key=lambda j: (hits[j], names[j] in _HEADERS["key"]))
        pick("key", i, f"{hits[i]} valeur(s) correspondent à des findings importés")
    v_hits = [sum(1 for v in col if guess_verdict(v)) for col in columns]
    v_order = sorted(
        (j for j in range(len(headers)) if j not in taken),
        key=lambda j: (names[j] in _HEADERS["verdict"], v_hits[j]),
        reverse=True,
    )
    if v_order and (names[v_order[0]] in _HEADERS["verdict"] or v_hits[v_order[0]]):
        j = v_order[0]
        pick("verdict", j, "en-tête reconnu" if names[j] in _HEADERS["verdict"] else f"{v_hits[j]} verdict(s) reconnus")
    for role in ("comment", "category", "line", "minutes", "file"):
        j = next((j for j, n in enumerate(names) if j not in taken and n in _HEADERS[role]), None)
        if j is not None:
            pick(role, j, "en-tête reconnu")
    if mapping["file"] is None:
        for j, col in enumerate(columns):
            texts = [t for t in (_text(v) for v in col) if t]
            if j in taken or not texts:
                continue
            same = sum(1 for t in texts if t.replace("\\", "/").split("/")[-1].split(":")[0].lower() in basenames)
            if same and same >= len(texts) / 2:
                pick("file", j, f"{same} nom(s) de fichier connus")
                break
    if mapping["key"] is None and mapping["file"] is None:
        key_col = next((j for j, n in enumerate(names) if n in _HEADERS["key"]), None)
        if key_col is not None:
            pick("key", key_col, "en-tête reconnu (aucune valeur ne correspond encore à un finding)")
    return mapping, reasons


def _saved_profile(conn: sqlite3.Connection, sig: str) -> dict[str, Any] | None:
    row = conn.execute("SELECT * FROM manual_workbook_profile WHERE signature = ?", (sig,)).fetchone()
    if row is None:
        return None
    return {"mapping": loads(row["mapping_json"], {}), "value_map": loads(row["value_map_json"], {})}


def scan(
    conn: sqlite3.Connection,
    campaign_id: str,
    path: Path,
    *,
    sheet: str | None = None,
    tool_label: str | None = None,
    mapping: dict[str, int | None] | None = None,
    value_map: dict[str, str] | None = None,
) -> ManualPlan:
    """Lit le classeur, infère (ou applique) la lecture des colonnes et rapproche chaque ligne d'un finding."""
    findings = _findings(conn, campaign_id, tool_label)
    if not findings:
        raise ManualImportError("Aucun finding importé dans cette campagne : importer d'abord le rapport Fortify.")
    names, chosen, rows = _read_sheet(path, sheet)
    if not sheet:
        # Onglet le plus probable : celui dont le plus de lignes se rapprochent d'un finding.
        best = max(
            range(len(names)),
            key=lambda n: len(
                _plan(conn, path, names[n], names, rows[n], findings, None, None).by_status(
                    "importable", "identique", "conflit", "commentaire seul", "non reconnu"
                )
            ),
        )
        chosen, rows = names[best], rows[best]
    return _plan(conn, path, chosen, names, rows, findings, mapping, value_map)


def _plan(conn, path, sheet, names, rows, findings, mapping, value_map) -> ManualPlan:
    rows = list(rows)
    h = _header_row(rows) if rows else 0
    headers = [str(_text(c) or f"(colonne {i + 1})") for i, c in enumerate(rows[h])] if rows else []
    body = rows[h + 1 :]
    sig = signature([x for x in headers if not x.startswith("(colonne")])
    saved = _saved_profile(conn, sig) if mapping is None else None
    reasons: dict[str, str] = {}
    if mapping is None and saved:
        by_name = {norm_name(x): i for i, x in enumerate(headers)}
        mapping = {r: by_name.get(saved["mapping"].get(r) or "") for r in ROLES}
        reasons = {r: "lecture enregistrée lors d'un import précédent" for r, i in mapping.items() if i is not None}
    elif mapping is None:
        mapping, reasons = _infer(headers, body, findings)
    else:
        mapping = {r: (i if i is not None and 0 <= i < len(headers) else None) for r, i in mapping.items()}
        reasons = {r: "choisi dans l'aperçu" for r, i in mapping.items() if i is not None}
    mapping = {r: mapping.get(r) for r in ROLES}
    plan = ManualPlan(
        path,
        sheet,
        names,
        h + 1,
        headers,
        mapping,
        reasons,
        {},
        Counter(),
        sig,
        profile_reused=bool(saved),
    )

    def cell(row: tuple, role: str) -> Any:
        i = mapping.get(role)
        return row[i] if i is not None and i < len(row) else None

    for n, row in enumerate(body, start=h + 2):
        raw = _text(cell(row, "verdict"))
        comment = _text(cell(row, "comment"))
        if raw is None and comment is None:
            continue  # ligne pas encore analysée
        file_v, line_v = _text(cell(row, "file")), cell(row, "line")
        line = apply_transform("int", line_v) if line_v not in (None, "") else None
        if file_v and line is None:
            line = apply_transform("location_line", file_v)
        if file_v:
            file_v = apply_transform("location_path", file_v)
        item = ManualItem(
            n,
            _text(cell(row, "key")),
            file_v,
            line,
            raw,
            comment,
            _number(cell(row, "minutes")),
            _text(cell(row, "category")),
        )
        if raw is not None:
            plan.value_counts[norm_name(raw)] += 1
        plan.items.append(item)
    saved_values = (saved or {}).get("value_map", {})
    for value in plan.value_counts:
        given = (value_map or {}).get(value)
        plan.value_map[value] = given or saved_values.get(value) or _VERDICT_GUESSES.get(value) or IGNORE
    seen: dict[str, int] = {}
    for item in plan.items:
        _match(item, findings)
        if not item.finding_id:
            continue
        if item.finding_id in seen:
            item.status, item.detail = "doublon", f"même finding que la ligne {seen[item.finding_id]}"
            continue
        seen[item.finding_id] = item.row
        _classify(conn, plan, item)
    return plan


def _match(item: ManualItem, findings: list[dict[str, Any]]) -> None:
    if item.key:
        found = [f for f in findings if f["source_id"].lower() == item.key.lower()]
        if len(found) == 1:
            item.finding_id, item.source_id, item.matched_by = found[0]["id"], found[0]["source_id"], "identifiant"
            return
        if len(found) > 1:
            item.status, item.detail = "ambigu", f"{len(found)} findings portent cet identifiant (outils différents)"
            return
    path = _norm_path(item.file)
    if path and item.line is not None:
        found = [
            f
            for f in findings
            if f["line_number"] == item.line
            and any(
                _same_file(path, p) for p in {_norm_path(f["normalized_path"]), _norm_path(f["full_filename"])} if p
            )
        ]
        if len(found) > 1 and item.category:
            narrowed = [f for f in found if f["category"] and norm_name(f["category"]) == norm_name(item.category)]
            found = narrowed or found
        if len(found) == 1:
            item.finding_id, item.source_id, item.matched_by = found[0]["id"], found[0]["source_id"], "fichier + ligne"
            return
        if len(found) > 1:
            ids = ", ".join(f["source_id"][-8:] for f in found[:4])
            item.status, item.detail = "ambigu", f"{len(found)} findings à {item.file}:{item.line} ({ids}…)"
            return
    item.status = "sans finding"
    item.detail = (
        "aucun finding importé avec cet identifiant" if item.key else "ni identifiant ni fichier + ligne connus"
    )


def _classify(conn: sqlite3.Connection, plan: ManualPlan, item: ManualItem) -> None:
    if item.raw_result is None:
        item.status, item.detail = "commentaire seul", "pas de verdict : non repris (rien n'est inventé)"
        return
    target = plan.value_map.get(norm_name(item.raw_result), IGNORE)
    if target == IGNORE:
        item.status = "non reconnu"
        item.detail = f"« {item.raw_result} » : traduction à choisir dans l'aperçu (True Positive / Not an issue)"
        return
    item.verdict = target
    f = conn.execute("SELECT current_decision_id FROM finding WHERE id = ?", (item.finding_id,)).fetchone()
    if f["current_decision_id"]:
        cur = dict(conn.execute("SELECT * FROM decision_event WHERE id = ?", (f["current_decision_id"],)).fetchone())
        if excel_projection(cur) == (verdict_to_excel(Verdict(target)), item.comment):
            item.status, item.detail = "identique", "décision Paladin déjà identique"
        else:
            p_result, p_comment = excel_projection(cur)
            item.status = "conflit"
            item.detail = (
                f"Paladin : {p_result or 'à investiguer'}{f' « {p_comment} »' if p_comment else ''} — conservé ;"
                " corriger dans Paladin si votre classeur est plus juste"
            )
        return
    item.status = "importable"
    if norm_name(item.raw_result) not in ("true positive", "not an issue"):
        item.detail = f"« {item.raw_result} » → « {verdict_to_excel(Verdict(target))} »"


@dataclass
class ManualImportResult:
    batch_id: str
    imported: list[str]
    roles: dict[str, int]
    baseline: dict[str, Any] | None


def apply(
    conn: sqlite3.Connection,
    campaign_id: str,
    plan: ManualPlan,
    author: str,
    *,
    role: str = "auto",
) -> ManualImportResult:
    """Enregistre les lignes importables comme décisions « import » ; mémorise la lecture du classeur.

    `role` : « auto » répartit exemples / référence (calibration), « example » ou « reference » les fixe tous.
    """
    from paladin import calibration

    if role not in ("auto", "example", "reference"):
        raise ManualImportError(f"Rôle inconnu : {role}")
    if not plan.importable:
        raise ManualImportError("Rien à importer : aucune ligne rapprochée d'un finding avec un verdict traduit.")
    batch_id = new_id()
    with store.transaction(conn):
        conn.execute(
            "INSERT INTO batch (id, group_id, member_ids_json, excluded_json, verdict, comment, author, created_at,"
            " campaign_id) VALUES (?, NULL, ?, ?, 'IMPORT', ?, ?, ?, ?)",
            (
                batch_id,
                dumps(sorted(i.finding_id for i in plan.importable)),
                dumps(
                    [
                        {"source_id": i.source_id or i.key or f"ligne {i.row}", "reason": f"{i.status} — {i.detail}"}
                        for i in plan.items
                        if i.status != "importable"
                    ]
                ),
                f"analyses manuelles : {plan.path.name} / {plan.sheet}",
                author,
                utcnow(),
                campaign_id,
            ),
        )
        conn.execute(
            "INSERT INTO manual_workbook_profile (signature, mapping_json, value_map_json, updated_at)"
            " VALUES (?, ?, ?, ?) ON CONFLICT (signature) DO UPDATE SET mapping_json = excluded.mapping_json,"
            " value_map_json = excluded.value_map_json, updated_at = excluded.updated_at",
            (
                plan.signature,
                dumps({r: norm_name(plan.headers[i]) if i is not None else None for r, i in plan.mapping.items()}),
                dumps(plan.value_map),
                utcnow(),
            ),
        )
    imported = []
    for item in plan.importable:
        f = conn.execute("SELECT revision FROM finding WHERE id = ?", (item.finding_id,)).fetchone()
        decisions.record_decision(
            conn,
            item.finding_id,
            expected_revision=f["revision"],
            action=DecisionAction.IMPORT,
            author=f"{author} (analyse manuelle)",
            verdict=item.verdict,
            comment=item.comment,
            authority=Authority.IMPORT,
            batch_id=batch_id,
        )
        imported.append(item.finding_id)
    roles = calibration.assign_roles(conn, campaign_id, imported, role)
    baseline = None
    if plan.minutes_total is not None:
        counted = [i for i in plan.items if i.minutes is not None and i.status in ("importable", "identique")]
        baseline = calibration.set_baseline(
            conn, campaign_id, plan.minutes_total, len(counted), source=f"colonne « {plan.column_label('minutes')} »"
        )
    return ManualImportResult(batch_id, imported, roles, baseline)


def stored_copy(campaign_dir: Path, name: str) -> Path:
    """Copie conservée d'un classeur déposé dans l'interface ; refuse tout chemin hors du dossier."""
    folder = (campaign_dir / "manual").resolve()
    path = (folder / Path(name).name).resolve()
    if path.parent != folder or not path.is_file():
        raise ManualImportError("Classeur déposé introuvable : le déposer à nouveau.")
    return path


def save_upload(campaign_dir: Path, filename: str, data: bytes) -> Path:
    folder = campaign_dir / "manual"
    folder.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(filename or "classeur.xlsx").name)[-80:] or "classeur.xlsx"
    path = folder / f"{utcnow()[:19].replace(':', '').replace('-', '')}_{safe}"
    path.write_bytes(data)
    return path
