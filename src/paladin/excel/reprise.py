"""Reprise d'un classeur déjà renseigné (analyse antérieure, saisie manuelle).

Les verdicts et commentaires déjà présents dans `analysis result` /
`Analysis result comment` deviennent des décisions Paladin traçables : action et
autorité « import », regroupées sous un identifiant de reprise annulable en bloc.
Rien n'est écrasé : une décision Paladin existante et différente est listée en
conflit et conservée. Une valeur non standard (`TP`, `FP`…) est normalisée et
signalée ; elle sera réécrite au texte exact au prochain export.
"""

from __future__ import annotations

import sqlite3
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from paladin import store
from paladin.config import Settings
from paladin.contracts import (
    EXCEL_ANALYSIS_RESULT_VALUES,
    Authority,
    DecisionAction,
    ExportState,
    Verdict,
    normalize_verdict,
    verdict_to_excel,
)
from paladin.excel.export import _key_text, _norm_header, sheet_plans, target_path
from paladin.review import decisions
from paladin.review.decisions import excel_projection
from paladin.util import dumps, new_id, sha256_file, utcnow


class RepriseError(ValueError):
    pass


@dataclass
class RepriseItem:
    sheet: str
    row: int
    key: str
    finding_id: str | None
    raw_result: Any
    comment: str | None
    verdict: str | None
    status: str  # importable | normalisé | identique | conflit | commentaire seul | non reconnu | sans finding
    detail: str = ""


@dataclass
class ReprisePlan:
    workbook: Path
    items: list[RepriseItem] = field(default_factory=list)

    def by_status(self, *statuses: str) -> list[RepriseItem]:
        return [i for i in self.items if i.status in statuses]

    @property
    def importable(self) -> list[RepriseItem]:
        return self.by_status("importable", "normalisé")


def scan(settings: Settings, conn: sqlite3.Connection, campaign_id: str) -> ReprisePlan:
    campaign = store.get_campaign(conn, campaign_id)
    path = target_path(settings, campaign)
    if not path.is_file():
        raise RepriseError(f"Classeur cible introuvable : {path}")
    plan = ReprisePlan(path)
    plans, _ = sheet_plans(conn, campaign)
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        for sp in plans:
            if sp.sheet not in wb.sheetnames:
                continue
            ws = wb[sp.sheet]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            headers = {_norm_header(v): i for i, v in enumerate(rows[0]) if _norm_header(v)}
            key_i = headers.get(_norm_header(sp.key_header))
            res_i = headers.get("analysis result")
            com_i = headers.get("analysis result comment")
            if key_i is None or res_i is None:
                continue
            for n, row in enumerate(rows[1:], start=2):
                key = _key_text(row[key_i] if key_i < len(row) else None)
                raw = row[res_i] if res_i < len(row) else None
                comment = row[com_i] if com_i is not None and com_i < len(row) else None
                comment = str(comment).strip() if comment not in (None, "") else None
                if not key or (raw in (None, "") and comment is None):
                    continue
                plan.items.append(_classify(conn, sp.tool_id, sp.sheet, n, key, raw, comment))
    finally:
        wb.close()
    return plan


def _classify(conn, tool_id, sheet, row, key, raw, comment) -> RepriseItem:
    f = conn.execute("SELECT * FROM finding WHERE tool_id = ? AND source_id = ?", (tool_id, key)).fetchone()
    item = RepriseItem(sheet, row, key, f["id"] if f else None, raw, comment, None, "")
    if f is None:
        item.status, item.detail = "sans finding", "aucun finding importé avec cette clé"
        return item
    text = str(raw).strip() if raw not in (None, "") else ""
    if not text:
        item.status, item.detail = "commentaire seul", "pas de verdict : non repris (rien n'est inventé)"
        return item
    if text in EXCEL_ANALYSIS_RESULT_VALUES:
        verdict = Verdict.TRUE_POSITIVE if text == verdict_to_excel(Verdict.TRUE_POSITIVE) else Verdict.NOT_AN_ISSUE
        status = "importable"
    else:
        try:
            verdict = normalize_verdict(text)
        except ValueError:
            verdict = None
        if verdict not in (Verdict.TRUE_POSITIVE, Verdict.NOT_AN_ISSUE):
            item.status, item.detail = "non reconnu", f"« {text} » n'est ni True Positive ni Not an issue"
            return item
        status = "normalisé"
        item.detail = f"« {text} » → « {verdict_to_excel(verdict)} » (réécrit au prochain export)"
    item.verdict = verdict.value
    if f["current_decision_id"]:
        cur = dict(conn.execute("SELECT * FROM decision_event WHERE id = ?", (f["current_decision_id"],)).fetchone())
        if excel_projection(cur) == (verdict_to_excel(verdict), comment):
            item.status, item.detail = "identique", "décision Paladin déjà identique"
        else:
            p_result, p_comment = excel_projection(cur)
            item.status = "conflit"
            item.detail = (
                f"Paladin : {p_result or 'à investiguer'}{f' « {p_comment} »' if p_comment else ''} — conservé"
            )
        return item
    item.status = status
    return item


def apply(settings: Settings, conn: sqlite3.Connection, campaign_id: str, author: str) -> tuple[str, ReprisePlan]:
    """Importe les éléments repris ; retourne (identifiant de reprise, plan)."""
    plan = scan(settings, conn, campaign_id)
    if not plan.importable:
        raise RepriseError("Rien à reprendre : aucune valeur importable dans le classeur.")
    reprise_id = new_id()
    now = utcnow()
    run_id = new_id()
    with store.transaction(conn):
        conn.execute(
            "INSERT INTO batch (id, group_id, member_ids_json, excluded_json, verdict, comment, author, created_at,"
            " campaign_id) VALUES (?, NULL, ?, ?, 'REPRISE', ?, ?, ?, ?)",
            (
                reprise_id,
                dumps(sorted(i.finding_id for i in plan.importable)),
                dumps(
                    [
                        {"source_id": i.key, "reason": f"{i.status} — {i.detail}"}
                        for i in plan.items
                        if i not in plan.importable
                    ]
                ),
                f"reprise du classeur {plan.workbook.name}",
                author,
                now,
                campaign_id,
            ),
        )
        # Les cellules reprises appartiennent désormais à Paladin : un export ultérieur pourra les normaliser.
        conn.execute(
            "INSERT INTO export_run (id, campaign_id, mode, source_path, source_sha256, destination_path,"
            " destination_sha256, status, summary_json, created_at, finished_at)"
            " VALUES (?, ?, 'reprise', ?, ?, ?, ?, 'verified', ?, ?, ?)",
            (
                run_id,
                campaign_id,
                str(plan.workbook),
                sha256_file(plan.workbook),
                str(plan.workbook),
                sha256_file(plan.workbook),
                dumps({"reprise_id": reprise_id}),
                now,
                now,
            ),
        )
    for item in plan.importable:
        f = conn.execute("SELECT revision FROM finding WHERE id = ?", (item.finding_id,)).fetchone()
        ev = decisions.record_decision(
            conn,
            item.finding_id,
            expected_revision=f["revision"],
            action=DecisionAction.IMPORT,
            author=f"{author} (reprise Excel)",
            verdict=item.verdict,
            comment=item.comment,
            authority=Authority.IMPORT,
            batch_id=reprise_id,
        )
        for key, value in (("analyst_result", item.raw_result), ("analyst_comment", item.comment)):
            if value not in (None, ""):
                conn.execute(
                    "INSERT OR IGNORE INTO export_cell (export_run_id, finding_id, sheet_name, cell_ref, column_key,"
                    " old_value, new_value, decision_event_id) VALUES (?, ?, ?, ?, ?, NULL, ?, ?)",
                    (run_id, item.finding_id, item.sheet, f"{key}:{item.row}", key, dumps(value), ev["id"]),
                )
        exact = item.status == "importable"
        conn.execute(
            "UPDATE finding SET export_state = ?, exported_decision_id = ?, exported_run_id = ? WHERE id = ?",
            (
                ExportState.EXPORTED.value if exact else ExportState.STALE.value,
                ev["id"] if exact else None,
                run_id,
                item.finding_id,
            ),
        )
    return reprise_id, plan


def change_target(settings: Settings, conn: sqlite3.Connection, campaign_id: str, path: Path) -> list[str]:
    """Change le classeur cible d'une campagne. Retourne les onglets attendus absents (avertissement)."""
    path = path.expanduser().resolve()
    if not path.is_file() or path.suffix.lower() != ".xlsx":
        raise RepriseError(f"Classeur .xlsx introuvable : {path}")
    campaign = store.get_campaign(conn, campaign_id)
    wb = load_workbook(path, read_only=True)
    try:
        names = set(wb.sheetnames)
    finally:
        wb.close()
    missing = [
        t["sheet_name"] for t in store.list_tools(conn, campaign_id) if t["sheet_name"] and t["sheet_name"] not in names
    ]
    cfg = campaign["config"]
    cfg["target_workbook"] = str(path)
    store.update_campaign_config(conn, campaign_id, cfg)
    conn.execute("DELETE FROM target_workbook WHERE campaign_id = ?", (campaign_id,))
    return missing
