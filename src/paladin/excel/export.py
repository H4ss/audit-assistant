"""Export Excel fiable (section 14).

Garanties :
  * écriture uniquement dans les colonnes autorisées (analyste ; métadonnées
    seulement pour les lignes générées en mode `generate_rows`) ;
  * rapprochement par clé à chaque export : un classeur réordonné reste sûr, une
    clé en double bloque les lignes concernées ;
  * une valeur analyste saisie hors Paladin n'est jamais écrasée sans
    autorisation explicite ;
  * texte écrit comme texte, jamais comme formule ;
  * fichier temporaire, vérification par réouverture (cellules ciblées et
    toutes les autres cellules inchangées), sauvegarde, remplacement atomique ;
  * fichier verrouillé ou erreur : aucune décision marquée exportée ;
  * manifeste reliant chaque cellule modifiée à sa décision.
"""

from __future__ import annotations

import json
import shutil
import sqlite3
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from paladin import store
from paladin.config import Settings
from paladin.contracts import (
    DEFAULT_COLUMNS,
    HEADER_ALIASES,
    ExportState,
)
from paladin.excel.sheets import current_schema, finding_value
from paladin.matching import Projection, projection
from paladin.review.decisions import excel_projection
from paladin.util import dumps, file_stamp, loads, new_id, sha256_file, utcnow

ANALYST_COLUMNS = {"analysis result": "analyst_result", "analysis result comment": "analyst_comment"}
_DEFAULT_BY_HEADER = {c.header.strip().lower(): c for c in DEFAULT_COLUMNS}
_DEFAULT_BY_HEADER.update({alias.lower(): c for c in DEFAULT_COLUMNS for alias in HEADER_ALIASES.get(c.key, ())})
TEXT_KEYS = {c.key for c in DEFAULT_COLUMNS if c.text} | {"source_comments", "category", "full_filename"}

# Parties OOXML que openpyxl ne préserve pas : classeur non qualifié pour l'écriture.
UNSUPPORTED_PARTS = {
    "xl/vbaProject.bin": "macros VBA",
    "xl/drawings/": "dessins / images",
    "xl/charts/": "graphiques",
    "xl/pivotTables/": "tableaux croisés dynamiques",
    "xl/pivotCache/": "caches de tableaux croisés",
    "xl/embeddings/": "objets incorporés",
    "_xmlsignatures/": "signatures numériques",
    "xl/slicers/": "segments",
}


class ExportBlockedError(RuntimeError):
    def __init__(self, message: str, action: str, status: str = "failed") -> None:
        super().__init__(message)
        self.action = action
        self.status = status


@dataclass
class SheetPlan:
    tool_id: str
    tool_label: str
    sheet: str
    mode: str  # complete_existing | generate_rows
    key_header: str
    key_field: str  # attribut du finding servant de clé (source_id)
    schema_columns: list[dict[str, Any]] | None = None  # schéma validé (onglet créé par Paladin)
    extra_headers: list[str] = field(default_factory=list)  # colonnes comparatives ajoutées sur validation


@dataclass
class CellWrite:
    sheet: str
    row: int
    col: int
    column_key: str
    old: Any
    new: Any
    finding_id: str | None
    decision_event_id: str | None
    text: bool
    relation_ids: list[str] | None = None

    @property
    def ref(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"


@dataclass
class ExportResult:
    run_id: str
    status: str
    destination: Path
    manifest_path: Path | None = None
    backup_path: Path | None = None
    summary: dict[str, Any] = field(default_factory=dict)
    error: str | None = None
    action: str | None = None


# ---------------------------------------------------------------------------
# Plans d'onglets
# ---------------------------------------------------------------------------


def sheet_plans(conn: sqlite3.Connection, campaign: dict[str, Any]) -> tuple[list[SheetPlan], list[str]]:
    """Onglets à écrire et outils sans onglet (qui restent hors export)."""
    plans, without_sheet = [], []
    for spec in campaign["config"].get("tools", []):
        tool = store.get_tool(conn, campaign["id"], spec["label"])
        if not tool["sheet_name"]:
            without_sheet.append(spec["label"])
            continue
        sheet_cfg = spec.get("sheet") or {}
        extensions = (campaign["config"].get("sheet_extensions") or {}).get(tool["sheet_name"], [])
        schema = current_schema(conn, tool["id"])
        if schema is not None:
            cols = loads(schema["columns_json"], [])
            key_header = next(c["header"] for c in cols if c["role"] == "key")
            plans.append(
                SheetPlan(
                    tool["id"],
                    spec["label"],
                    tool["sheet_name"],
                    "generate_rows",
                    key_header,
                    "source_id",
                    cols,
                    extensions,
                )
            )
            continue
        if spec["kind"] == "fortify":
            mode = sheet_cfg.get("mode", "generate_rows")
            key_header = sheet_cfg.get("key_header", "Instance ID")
        else:
            mode = sheet_cfg.get("mode", "complete_existing")
            inventory = next((s for s in spec["sources"] if s["role"] == "inventory"), None)
            default_key = (inventory or {}).get("mapping", {}).get("fields", {}).get("source_id", {}).get("source")
            key_header = sheet_cfg.get("key_header", default_key or "ID")
        plans.append(
            SheetPlan(tool["id"], spec["label"], tool["sheet_name"], mode, key_header, "source_id", None, extensions)
        )
    return plans, without_sheet


# ---------------------------------------------------------------------------
# Contrôles préalables
# ---------------------------------------------------------------------------


def qualify_workbook(path: Path) -> None:
    if path.suffix.lower() != ".xlsx":
        raise ExportBlockedError(
            f"Format non qualifié : {path.suffix}", "Le MVP écrit uniquement des classeurs .xlsx standard."
        )
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
    except zipfile.BadZipFile as exc:
        raise ExportBlockedError(f"Classeur illisible : {exc}", "Vérifier que le fichier est un .xlsx valide.") from exc
    found = sorted({label for part, label in UNSUPPORTED_PARTS.items() for n in names if n.startswith(part)})
    if found:
        raise ExportBlockedError(
            f"Classeur contenant des éléments non préservés à l'écriture : {', '.join(found)}.",
            "Qualification séparée requise. Exporter vers un classeur sans ces éléments ou les retirer de la cible.",
        )


def lock_markers(path: Path) -> list[Path]:
    """Fichiers de verrou laissés par Excel (~$nom) ou LibreOffice (.~lock.nom#)."""
    candidates = [
        path.with_name("~$" + path.name),
        path.with_name("~$" + path.name[2:]),
        path.with_name(f".~lock.{path.name}#"),
    ]
    return [p for p in candidates if p.exists()]


# ---------------------------------------------------------------------------
# Lecture de l'état actuel d'un onglet
# ---------------------------------------------------------------------------


def _norm_header(v: Any) -> str:
    return str(v).strip().lower() if v is not None else ""


def _cell_value(v: Any) -> Any:
    if isinstance(v, str):
        return v if v.strip() else None
    return v


def _header_map(ws: Worksheet, header_row: int = 1) -> dict[str, int]:
    out: dict[str, int] = {}
    for cell in ws[header_row]:
        h = _norm_header(cell.value)
        if h:
            out.setdefault(h, cell.column)
    return out


def _key_text(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def _last_written(conn: sqlite3.Connection, finding_id: str, column_key: str) -> tuple[bool, Any]:
    """Dernière valeur écrite par Paladin pour (finding, colonne), lors d'un export réussi."""
    row = conn.execute(
        "SELECT c.new_value FROM export_cell c JOIN export_run r ON r.id = c.export_run_id"
        " WHERE c.finding_id = ? AND c.column_key = ? AND r.status = 'verified'"
        " ORDER BY r.created_at DESC LIMIT 1",
        (finding_id, column_key),
    ).fetchone()
    if row is None:
        return False, None
    return True, loads(row["new_value"])


# ---------------------------------------------------------------------------
# Calcul des écritures
# ---------------------------------------------------------------------------


@dataclass
class _Plan:
    writes: list[CellWrite] = field(default_factory=list)
    appended_rows: dict[str, list[int]] = field(default_factory=dict)
    up_to_date: set[str] = field(default_factory=set)  # findings dont les cellules analyste sont exactes
    blocked: list[dict[str, Any]] = field(default_factory=list)
    without_target: list[dict[str, Any]] = field(default_factory=list)
    unmatched_rows: list[dict[str, Any]] = field(default_factory=list)
    preserved: list[dict[str, Any]] = field(default_factory=list)


def _metadata_columns(plan: SheetPlan, headers: dict[str, int]) -> dict[int, tuple[str, str, bool]]:
    """Colonnes de métadonnées écrites dans les lignes générées : {colonne: (clé, source, texte)}."""
    out: dict[int, tuple[str, str, bool]] = {}
    if plan.schema_columns is not None:
        for c in plan.schema_columns:
            col = headers.get(_norm_header(c["header"]))
            if col is None or c["role"] not in ("key", "metadata"):
                continue
            source = c["source"] if (c.get("source") or "").startswith("extra:") else c["key"]
            out[col] = (c["key"], source, c["key"] != "line_number")
        return out
    for h, col in headers.items():
        col_def = _DEFAULT_BY_HEADER.get(h)
        if col_def is not None and not col_def.analyst:
            out[col] = (col_def.key, col_def.key, col_def.text or col_def.key in TEXT_KEYS)
    return out


def _comparative_columns(
    conn: sqlite3.Connection, plan: SheetPlan, headers: dict[str, int]
) -> dict[int, tuple[str, str, str]]:
    """`Found in X` / `criticality in X` : {colonne: (type, id de l'outil X, libellé X)}."""
    tools = {
        t["label"].lower(): t
        for t in conn.execute(
            "SELECT * FROM tool WHERE campaign_id = (SELECT campaign_id FROM tool WHERE id = ?)", (plan.tool_id,)
        )
    }
    out: dict[int, tuple[str, str, str]] = {}
    for h, col in headers.items():
        for prefix, kind in (("found in ", "found_in"), ("criticality in ", "criticality_in")):
            if h.startswith(prefix):
                other = tools.get(h[len(prefix) :].strip())
                if other is not None and other["id"] != plan.tool_id:
                    out[col] = (kind, other["id"], other["label"])
    return out


def _plan_sheet(
    conn: sqlite3.Connection,
    ws: Worksheet,
    plan: SheetPlan,
    allow_overwrite: set[tuple[str, str]],
    out: _Plan,
) -> None:
    headers = _header_map(ws)
    key_col = headers.get(_norm_header(plan.key_header))
    if key_col is None:
        raise ExportBlockedError(
            f"Onglet {plan.sheet!r} : colonne clé {plan.key_header!r} introuvable.",
            "Vérifier les en-têtes de l'onglet ou le mapping de la campagne.",
        )
    analyst_cols = {key: headers[h] for h, key in ANALYST_COLUMNS.items() if h in headers}
    comparative_cols = _comparative_columns(conn, plan, headers)
    meta_cols = _metadata_columns(plan, headers)
    projections: dict[tuple[str, str], Projection] = {}
    missing = [h for h in ANALYST_COLUMNS if h not in headers]
    if missing:
        raise ExportBlockedError(
            f"Onglet {plan.sheet!r} : colonnes analyste absentes : {missing}.",
            "Ajouter les colonnes `analysis result` et `Analysis result comment` ou corriger le mapping.",
        )

    index: dict[str, list[int]] = {}
    for r in range(2, ws.max_row + 1):
        k = _key_text(ws.cell(row=r, column=key_col).value)
        if k:
            index.setdefault(k, []).append(r)

    findings = conn.execute(
        "SELECT f.*, (SELECT COUNT(*) FROM fingerprint_collision c WHERE c.finding_id = f.id AND c.status = 'open')"
        " AS open_collisions FROM finding f WHERE f.tool_id = ? ORDER BY f.created_at, f.source_id",
        (plan.tool_id,),
    ).fetchall()
    known_keys = set()
    next_row = ws.max_row + 1
    for f in findings:
        key = f[plan.key_field]
        known_keys.add(key)
        label = {"finding_id": f["id"], "source_id": f["source_id"], "sheet": plan.sheet}
        if f["open_collisions"]:
            out.blocked.append({**label, "reason": "collision d'identifiant non résolue"})
            continue
        rows = index.get(key, [])
        if len(rows) > 1:
            out.blocked.append({**label, "reason": f"clé présente sur plusieurs lignes : {rows}"})
            continue
        decision = None
        if f["current_decision_id"]:
            decision = dict(
                conn.execute("SELECT * FROM decision_event WHERE id = ?", (f["current_decision_id"],)).fetchone()
            )
        result, comment = excel_projection(decision)
        desired = {"analyst_result": result, "analyst_comment": comment}

        if not rows:
            match_state = (loads(f["details_json"], {}).get("match") or {}).get("state")
            if plan.mode != "generate_rows":
                reason = (
                    "MD sans ligne Excel : proposition de nouvelle ligne (politique d'ajout non validée)"
                    if match_state == "details_only"
                    else "aucune ligne avec cette clé"
                )
                out.without_target.append({**label, "reason": reason})
                continue
            row = next_row
            next_row += 1
            out.appended_rows.setdefault(plan.sheet, []).append(row)
            for col, (key, source, text) in meta_cols.items():
                value = finding_value(f, source)
                if value is not None:
                    out.writes.append(CellWrite(plan.sheet, row, col, key, None, value, f["id"], None, text))
            for col, (kind, other_id, other_label) in comparative_cols.items():
                proj = projections.setdefault((f["id"], other_id), projection(conn, f, other_id))
                value = proj.found_in if kind == "found_in" else proj.criticality
                if value is not None:
                    out.writes.append(
                        CellWrite(
                            plan.sheet,
                            row,
                            col,
                            f"{kind}:{other_label}",
                            None,
                            value,
                            f["id"],
                            None,
                            True,
                            proj.relation_ids,
                        )
                    )
            for key, col in analyst_cols.items():
                if desired[key] is not None:
                    out.writes.append(
                        CellWrite(
                            plan.sheet, row, col, key, None, desired[key], f["id"], f["current_decision_id"], True
                        )
                    )
            if f["current_decision_id"]:
                out.up_to_date.add(f["id"])
            continue

        row = rows[0]
        finding_ok = True
        for key, col in analyst_cols.items():
            current = _cell_value(ws.cell(row=row, column=col).value)
            new = desired[key]
            if current == new:
                continue
            wrote_before, last = _last_written(conn, f["id"], key)
            paladin_owned = current is None or (wrote_before and current == last)
            if new is None and not paladin_owned:
                # Aucune décision Paladin (ou investigation) : la saisie humaine est conservée.
                out.preserved.append({**label, "cell": f"{get_column_letter(col)}{row}", "value": current})
                continue
            if paladin_owned or (f["id"], key) in allow_overwrite:
                out.writes.append(
                    CellWrite(plan.sheet, row, col, key, current, new, f["id"], f["current_decision_id"], True)
                )
            else:
                finding_ok = False
                out.blocked.append(
                    {
                        **label,
                        "cell": f"{get_column_letter(col)}{row}",
                        "reason": "valeur saisie hors Paladin, différente de la décision : revue explicite requise",
                        "excel_value": current,
                        "paladin_value": new,
                    }
                )
        for col, (kind, other_id, other_label) in comparative_cols.items():
            proj = projections.setdefault((f["id"], other_id), projection(conn, f, other_id))
            new = proj.found_in if kind == "found_in" else proj.criticality
            column_key = f"{kind}:{other_label}"
            current = _cell_value(ws.cell(row=row, column=col).value)
            if current == new:
                continue
            wrote_before, last = _last_written(conn, f["id"], column_key)
            if current is None or (wrote_before and current == last) or (f["id"], column_key) in allow_overwrite:
                out.writes.append(
                    CellWrite(plan.sheet, row, col, column_key, current, new, f["id"], None, True, proj.relation_ids)
                )
            elif new is not None:
                out.preserved.append({**label, "cell": f"{get_column_letter(col)}{row}", "value": current})
        if finding_ok and f["current_decision_id"]:
            out.up_to_date.add(f["id"])

    for k, rows in index.items():
        if k not in known_keys:
            for r in rows:
                out.unmatched_rows.append({"sheet": plan.sheet, "row": r, "key": k})


# ---------------------------------------------------------------------------
# Écriture et vérification
# ---------------------------------------------------------------------------


def _apply(wb, writes: list[CellWrite]) -> None:
    for w in writes:
        cell = wb[w.sheet].cell(row=w.row, column=w.col)
        cell.value = w.new
        if w.text and isinstance(w.new, str):
            cell.data_type = "s"  # jamais interprété comme formule


def _snapshot(path: Path) -> dict[str, dict[str, Any]]:
    wb = load_workbook(path, data_only=False)
    try:
        return {
            ws.title: {
                c.coordinate: (c.value, c.data_type) for row in ws.iter_rows() for c in row if c.value is not None
            }
            for ws in wb.worksheets
        }
    finally:
        wb.close()


def _verify(source_snapshot: dict, written: Path, writes: list[CellWrite], sheet_order: list[str]) -> list[str]:
    problems: list[str] = []
    after = _snapshot(written)
    if list(after) != sheet_order:
        problems.append(f"ordre des onglets modifié : {list(after)}")
    targeted = {(w.sheet, w.ref): w for w in writes}
    for sheet, cells in source_snapshot.items():
        for coord, val in cells.items():
            if (sheet, coord) in targeted:
                continue
            if after.get(sheet, {}).get(coord) != val:
                problems.append(f"{sheet}!{coord} modifiée hors périmètre")
    for sheet, cells in after.items():
        for coord in cells:
            if (sheet, coord) not in targeted and coord not in source_snapshot.get(sheet, {}):
                problems.append(f"{sheet}!{coord} ajoutée hors périmètre")
    for (sheet, coord), w in targeted.items():
        got = after.get(sheet, {}).get(coord)
        got_value = got[0] if got else None
        if _cell_value(got_value) != w.new:
            problems.append(f"{sheet}!{coord} : attendu {w.new!r}, lu {got_value!r}")
        elif w.text and isinstance(w.new, str) and got and got[1] != "s":
            problems.append(f"{sheet}!{coord} : texte attendu, type {got[1]!r}")
    return problems[:50]


def _record_run(conn, campaign_id, mode, source, destination, source_sha) -> str:
    rid = new_id()
    conn.execute(
        "INSERT INTO export_run (id, campaign_id, mode, source_path, source_sha256, destination_path, status,"
        " created_at)"
        " VALUES (?, ?, ?, ?, ?, ?, 'pending', ?)",
        (rid, campaign_id, mode, str(source), source_sha, str(destination), utcnow()),
    )
    return rid


def _fail(conn, run_id, status, error, action, destination, summary=None) -> ExportResult:
    conn.execute(
        "UPDATE export_run SET status = ?, error = ?, summary_json = ?, finished_at = ? WHERE id = ?",
        (status, f"{error} → {action}", dumps(summary or {}), utcnow(), run_id),
    )
    return ExportResult(run_id, status, destination, summary=summary or {}, error=error, action=action)


def target_path(settings: Settings, campaign: dict[str, Any]) -> Path:
    p = Path(campaign["config"]["target_workbook"])
    return p if p.is_absolute() else settings.campaign_dir(campaign["id"]) / "inputs" / p


def export_workbook(
    settings: Settings,
    conn: sqlite3.Connection,
    campaign_id: str,
    *,
    mode: str = "working_copy",
    destination: Path | None = None,
    allow_overwrite: set[tuple[str, str]] | None = None,
) -> ExportResult:
    """Exporte les décisions validées.

    `working_copy` : lit le classeur cible et écrit une copie de travail dans
    l'espace de campagne (le fichier de l'utilisateur n'est pas touché).
    `final` : met à jour le classeur cible désigné, avec sauvegarde.
    """
    if mode not in ("working_copy", "final"):
        raise ValueError(mode)
    campaign = store.get_campaign(conn, campaign_id)
    source = target_path(settings, campaign)
    exports = settings.campaign_dir(campaign_id) / "exports"
    if destination is None:
        destination = source if mode == "final" else exports / f"{source.stem}.paladin{source.suffix}"
    exports.mkdir(parents=True, exist_ok=True)
    source_sha = sha256_file(source) if source.exists() else None
    run_id = _record_run(conn, campaign_id, mode, source, destination, source_sha)

    try:
        if not source.exists():
            raise ExportBlockedError(
                f"Classeur cible introuvable : {source}", "Vérifier le chemin du classeur dans la campagne."
            )
        qualify_workbook(source)
        locks = lock_markers(destination) + (lock_markers(source) if mode == "final" else [])
        if locks:
            raise ExportBlockedError(
                f"Classeur ouvert dans un tableur ({', '.join(p.name for p in locks)}).",
                "Fermer le classeur dans Excel puis relancer l'export. Les décisions restent enregistrées.",
                status="locked",
            )
        tw = conn.execute(
            "SELECT * FROM target_workbook WHERE campaign_id = ? AND path = ?", (campaign_id, str(source))
        ).fetchone()
        externally_modified = bool(tw and tw["known_sha256"] and tw["known_sha256"] != source_sha)

        plans, without_sheet = sheet_plans(conn, campaign)
        wb = load_workbook(source, data_only=False)
        sheet_order = list(wb.sheetnames)
        plan = _Plan()
        expected_order = list(sheet_order)
        for sp in plans:
            if sp.sheet not in wb.sheetnames:
                if sp.schema_columns is None:
                    raise ExportBlockedError(
                        f"Onglet {sp.sheet!r} absent du classeur.",
                        "Valider un schéma d'onglet pour cet outil (page « Nouvel onglet »)"
                        " ou corriger la configuration.",
                    )
                ws_new = wb.create_sheet(sp.sheet)
                expected_order.append(sp.sheet)
                for i, c in enumerate(sp.schema_columns, start=1):
                    ws_new.cell(row=1, column=i, value=c["header"])
                    plan.writes.append(CellWrite(sp.sheet, 1, i, "header", None, c["header"], None, None, True))
            ws = wb[sp.sheet]
            existing = _header_map(ws)
            next_col = ws.max_column + 1 if any(c.value is not None for c in ws[1]) else 1
            for header in sp.extra_headers:
                if _norm_header(header) not in existing:
                    ws.cell(row=1, column=next_col, value=header)
                    plan.writes.append(CellWrite(sp.sheet, 1, next_col, "header", None, header, None, None, True))
                    existing[_norm_header(header)] = next_col
                    next_col += 1
            _plan_sheet(conn, ws, sp, allow_overwrite or set(), plan)
        _apply(wb, plan.writes)
        tmp = destination.with_name(f".~paladin-{run_id[:8]}{destination.suffix}")
        destination.parent.mkdir(parents=True, exist_ok=True)
        wb.save(tmp)
        wb.close()
        problems = _verify(_snapshot(source), tmp, plan.writes, expected_order)
        if problems:
            tmp.unlink(missing_ok=True)
            raise ExportBlockedError(
                "Vérification après écriture échouée : " + "; ".join(problems[:5]),
                "Aucun fichier remplacé. Signaler ce cas (classeur non qualifié ?).",
            )

        backup = None
        if destination.exists():
            backups = settings.campaign_dir(campaign_id) / "backups"
            backups.mkdir(parents=True, exist_ok=True)
            backup = backups / f"{destination.stem}.{file_stamp()}-{run_id[:6]}{destination.suffix}"
            shutil.copy2(destination, backup)
        try:
            tmp.replace(destination)
        except PermissionError as exc:
            tmp.unlink(missing_ok=True)
            raise ExportBlockedError(
                f"Remplacement impossible : {exc}",
                "Fermer le classeur dans Excel puis relancer l'export. Les décisions restent enregistrées.",
                status="locked",
            ) from exc
    except ExportBlockedError as exc:
        return _fail(conn, run_id, exc.status, str(exc), exc.action, destination)
    except PermissionError as exc:
        return _fail(
            conn,
            run_id,
            "locked",
            f"Accès refusé : {exc}",
            "Fermer le classeur dans Excel puis relancer l'export. Les décisions restent enregistrées.",
            destination,
        )

    dest_sha = sha256_file(destination)
    summary = {
        "cells_written": len(plan.writes),
        "rows_appended": {k: len(v) for k, v in plan.appended_rows.items()},
        "findings_up_to_date": len(plan.up_to_date),
        "blocked": plan.blocked,
        "without_target": plan.without_target,
        "unmatched_rows": plan.unmatched_rows,
        "preserved_human_values": plan.preserved,
        "tools_without_sheet": without_sheet,
        "externally_modified_since_last_export": externally_modified,
        "comparative_cells_written": sum(
            1 for w in plan.writes if w.column_key.startswith(("found_in:", "criticality_in:"))
        ),
        "sheets_created": [s for s in expected_order if s not in sheet_order],
    }
    manifest_path = exports / f"export-{run_id}.manifest.json"
    with store.transaction(conn):
        for w in plan.writes:
            conn.execute(
                "INSERT INTO export_cell (export_run_id, finding_id, sheet_name, cell_ref, column_key, old_value,"
                " new_value, decision_event_id, relation_ids_json) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    run_id,
                    w.finding_id,
                    w.sheet,
                    w.ref,
                    w.column_key,
                    dumps(w.old),
                    dumps(w.new),
                    w.decision_event_id,
                    dumps(w.relation_ids) if w.relation_ids else None,
                ),
            )
        for fid in plan.up_to_date:
            conn.execute(
                "UPDATE finding SET export_state = ?, exported_decision_id = current_decision_id, exported_run_id = ?"
                " WHERE id = ?",
                (ExportState.EXPORTED.value, run_id, fid),
            )
        conn.execute(
            "UPDATE export_run SET status = 'verified', destination_sha256 = ?, backup_path = ?, summary_json = ?,"
            " finished_at = ? WHERE id = ?",
            (dest_sha, str(backup) if backup else None, dumps(summary), utcnow(), run_id),
        )
        if mode == "final":
            conn.execute(
                "INSERT INTO target_workbook (id, campaign_id, path, known_sha256, updated_at) VALUES (?, ?, ?, ?, ?)"
                " ON CONFLICT (campaign_id, path) DO UPDATE SET known_sha256 = excluded.known_sha256,"
                " updated_at = excluded.updated_at",
                (new_id(), campaign_id, str(source), dest_sha, utcnow()),
            )
    manifest = {
        "run_id": run_id,
        "mode": mode,
        "source": str(source),
        "source_sha256": source_sha,
        "destination": str(destination),
        "destination_sha256": dest_sha,
        "backup": str(backup) if backup else None,
        "created_at": utcnow(),
        "summary": summary,
        "cells": [
            {
                "sheet": w.sheet,
                "cell": w.ref,
                "column": w.column_key,
                "old": w.old,
                "new": w.new,
                "finding_id": w.finding_id,
                "decision_event_id": w.decision_event_id,
            }
            for w in plan.writes
        ],
    }
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    return ExportResult(run_id, "verified", destination, manifest_path, backup, summary)
