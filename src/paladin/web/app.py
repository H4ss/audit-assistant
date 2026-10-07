"""Application web locale (FastAPI + Jinja2, sans CDN ni compilation).

Sécurité de l'interface locale :
  * écoute sur la boucle locale uniquement (vérifié au lancement) ;
  * en-tête Host restreint à la boucle locale (protection contre le DNS rebinding) ;
  * toute écriture exige le jeton d'interface (autorité humaine), distinct du
    jeton de l'agent, et une origine locale (protection CSRF).
"""

from __future__ import annotations

import shlex
import subprocess
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any
from urllib.parse import urlencode, urlparse

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

from paladin import __version__, store
from paladin.agent import jobs as agent_jobs
from paladin.analysis import safe_repo_file
from paladin.config import Settings
from paladin.contracts import (
    DISCUSSION_COMMENT,
    EXCEL_NOT_AN_ISSUE,
    EXCEL_TRUE_POSITIVE,
    DecisionAction,
    Verdict,
)
from paladin.db import open_database
from paladin.review import decisions as dec
from paladin.review import queue as q
from paladin.store import ConflictError
from paladin.util import loads, utcnow

_HERE = Path(__file__).parent
LOOPBACK_HOSTS = {"127.0.0.1", "localhost", "::1"}
AUTHOR = "analyste"

VERDICT_LABELS = {
    "TRUE_POSITIVE": EXCEL_TRUE_POSITIVE,
    "NOT_AN_ISSUE": EXCEL_NOT_AN_ISSUE,
    "NEEDS_REVIEW": "À investiguer / indéterminé",
}
EXPORT_LABELS = {
    "not_exported": "Excel : à exporter",
    "exported": "Excel : à jour",
    "stale": "Excel : PÉRIMÉ — réexporter",
    "conflict": "Excel : conflit",
}
REVIEW_LABELS = {
    "to_review": "À revoir",
    "investigating": "À investiguer",
    "validated": "Validé",
    "reexam_required": "Réexamen requis",
}


def _flash(conn, campaign_id: str, message: str, level: str = "info", details: Any = None) -> None:
    store.set_ui_state(conn, campaign_id, "flash", {"message": message, "level": level, "details": details})


def _pop_flash(conn, campaign_id: str) -> dict | None:
    flash = store.get_ui_state(conn, campaign_id, "flash")
    if flash:
        conn.execute("DELETE FROM ui_state WHERE campaign_id = ? AND key = 'flash'", (campaign_id,))
    return flash


def create_app(settings: Settings) -> FastAPI:
    conn = open_database(settings.db_path)
    ui_token = settings.ui_token()

    @asynccontextmanager
    async def lifespan(_app: FastAPI) -> AsyncIterator[None]:
        yield
        conn.close()

    app = FastAPI(
        title="Paladin", version=__version__, docs_url=None, redoc_url=None, openapi_url=None, lifespan=lifespan
    )
    app.state.settings = settings
    app.state.conn = conn
    templates = Jinja2Templates(directory=str(_HERE / "templates"))
    templates.env.globals.update(
        version=__version__,
        token=ui_token,
        VERDICT_LABELS=VERDICT_LABELS,
        EXPORT_LABELS=EXPORT_LABELS,
        REVIEW_LABELS=REVIEW_LABELS,
        VIEWS=q.VIEWS,
        DISCUSSION_COMMENT=DISCUSSION_COMMENT,
    )
    app.mount("/static", StaticFiles(directory=str(_HERE / "static")), name="static")
    from paladin.agent.api import build_router
    from paladin.agent.workspace import skill_version

    app.include_router(build_router(conn, settings.agent_token(), skill_version()))

    @app.middleware("http")
    async def local_only(request: Request, call_next):
        host = urlparse("//" + (request.headers.get("host") or "")).hostname or ""
        if host and host not in LOOPBACK_HOSTS and host != "testserver":
            return Response("Hôte refusé : Paladin n'est accessible que sur la boucle locale.", status_code=421)
        if request.method in ("POST", "PUT", "DELETE"):
            origin = request.headers.get("origin") or request.headers.get("referer")
            if origin:
                o_host = urlparse(origin).hostname or ""
                if o_host not in {"127.0.0.1", "localhost", "::1", "testserver"}:
                    return Response("Origine refusée.", status_code=403)
        return await call_next(request)

    async def _form(request: Request) -> dict[str, str]:
        form = await request.form()
        data = {k: str(v) for k, v in form.items()}
        if data.get("token") != ui_token:
            raise HTTPException(403, "Jeton d'interface invalide : recharger la page.")
        return data

    def _campaign(cid: str) -> dict[str, Any]:
        try:
            return store.get_campaign(conn, cid)
        except store.NotFoundError:
            raise HTTPException(404, "Campagne inconnue") from None

    def render(request: Request, name: str, **ctx: Any) -> HTMLResponse:
        return templates.TemplateResponse(request, name, ctx)

    # ------------------------------------------------------------------ pages

    global_flash: dict[str, Any] = {}

    def _gflash(message: str, level: str = "info", details: Any = None) -> None:
        global_flash.update({"message": message, "level": level, "details": details})

    def _pop_gflash() -> dict | None:
        out = dict(global_flash) or None
        global_flash.clear()
        return out

    @app.get("/agent", response_class=HTMLResponse)
    def agent_page(request: Request, detect: int = 0):
        from paladin.agent import connect
        from paladin.agent.workspace import DEFAULT_MODEL

        info = connect.detect(settings) if detect else None
        return render(
            request,
            "agent.html",
            info=info,
            current=settings.agent.get("model"),
            default=DEFAULT_MODEL,
            flash=_pop_gflash(),
            campaigns=store.list_campaigns(conn),
        )

    @app.post("/agent/connect")
    async def agent_connect(request: Request):
        from paladin.agent import connect

        form = await _form(request)
        model = (form.get("custom") or form.get("model") or "").strip()
        try:
            res = connect.smoke_test(settings, model)
            connect.record_smoke(settings, res)
        except ValueError as exc:
            _gflash(str(exc), "warn")
            return RedirectResponse("/agent", status_code=303)
        if res.ok:
            connect.save_model(settings, model)
            _gflash(
                f"Connecté : {model} a répondu « {res.answer} » en {res.seconds:.0f} s"
                f" (coût {res.cost_usd or 0:.4f} $). Modèle enregistré.",
                "ok",
                ["Étape suivante recommandée : « Sondage de sécurité » ci-dessous."],
            )
        else:
            _gflash(f"Échec avec {model} : {res.error}", "error", [f"→ {res.action}"])
        return RedirectResponse("/agent", status_code=303)

    @app.post("/agent/probe")
    async def agent_probe(request: Request):
        from paladin.agent.runner import AgentRunError, probe_agent

        await _form(request)
        try:
            res = probe_agent(settings)
        except AgentRunError as exc:
            _gflash(str(exc), "error", [f"→ {exc.action}"])
            return RedirectResponse("/agent", status_code=303)
        if res.ok:
            _gflash(
                "Sondage OK : l'agent n'a exécuté aucun outil hors paladin_* (shell, fichiers, réseau).",
                "ok",
                [f"Réponse du modèle : {res.model_text[:300]}"],
            )
        else:
            _gflash(
                f"ALERTE : outils interdits exécutés {res.forbidden}. Ne pas utiliser l'agent.",
                "error",
                [f"Journal : {res.log}"],
            )
        return RedirectResponse("/agent", status_code=303)

    # ------------------------------------------------- PC de travail : checklist, SSC

    @app.get("/travail", response_class=HTMLResponse)
    def travail(request: Request):
        from paladin import readiness

        items = readiness.checklist(settings, conn)
        sections: dict[str, list] = {}
        for it in items:
            sections.setdefault(it.section, []).append(it)
        done = sum(1 for it in items if it.status == readiness.OK)
        return render(
            request,
            "travail.html",
            sections=sections,
            done=done,
            total=len(items),
            fortify=settings.fortify,
            flash=_pop_gflash(),
            campaigns=store.list_campaigns(conn),
        )

    @app.post("/travail/ssc")
    async def travail_ssc(request: Request):
        from paladin.config import set_config_value
        from paladin.fortify.ssc import save_token

        form = await _form(request)
        url = form.get("url", "").strip()
        ca = form.get("ca_bundle", "").strip()
        set_config_value(settings.config_path, "fortify", "url", url)
        set_config_value(settings.config_path, "fortify", "ca_bundle", ca)
        settings.fortify.update({"url": url, "ca_bundle": ca})
        if form.get("ssc_token", "").strip():
            save_token(settings, form["ssc_token"])
        _gflash("Paramètres SSC enregistrés (jeton stocké hors du dépôt). Lancer le diagnostic.", "ok")
        return RedirectResponse("/travail#ssc", status_code=303)

    @app.post("/travail/check")
    async def travail_check(request: Request):
        from paladin import readiness
        from paladin.fortify.check import STATE_BLOCKED, run_check, write_report

        await _form(request)
        rep = run_check(settings)
        write_report(settings, rep)
        readiness.record(settings, "fortify_check", rep.state != STATE_BLOCKED, rep.state, report="/travail/rapport")
        details = [
            f"[{c.status.value}] {c.area}/{c.name} : {c.detail}"
            + (f" → {c.action}" if c.action and c.status.value != "OK" else "")
            for c in rep.checks
        ]
        _gflash(f"Diagnostic : {rep.state}", "error" if rep.state == STATE_BLOCKED else "ok", details)
        return RedirectResponse("/travail#ssc", status_code=303)

    @app.get("/travail/rapport", response_class=HTMLResponse)
    def travail_rapport(request: Request):
        reports = sorted((settings.home / "reports").glob("fortify-check-*.md"))
        text = reports[-1].read_text(encoding="utf-8") if reports else "Aucun rapport : lancer le diagnostic."
        return render(request, "report.html", text=text, path=str(reports[-1]) if reports else "")

    def _discovery_file():
        return settings.home / "run" / "ssc_discovery.json"

    @app.get("/fortify/discover", response_class=HTMLResponse)
    def fortify_discover_page(request: Request):
        import json as _json

        path = _discovery_file()
        cached = _json.loads(path.read_text(encoding="utf-8")) if path.exists() else None
        return render(
            request, "discover.html", cached=cached, flash=_pop_gflash(), campaigns=store.list_campaigns(conn)
        )

    @app.post("/fortify/discover")
    async def fortify_discover_run(request: Request):
        import json as _json

        from paladin.fortify.discovery import discover
        from paladin.fortify.ssc import make_client
        from paladin.importers.fortify import FortifyError

        form = await _form(request)
        try:
            groups = discover(make_client(settings), name_filter=form.get("filter") or None)
        except FortifyError as exc:
            _gflash(str(exc), "error", [f"→ {exc.action}"])
            return RedirectResponse("/fortify/discover", status_code=303)
        path = _discovery_file()
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            _json.dumps(
                {"at": utcnow(), "filter": form.get("filter") or "", "groups": [g.as_dict() for g in groups]},
                ensure_ascii=False,
                default=str,
            ),
            encoding="utf-8",
        )
        return RedirectResponse("/fortify/discover", status_code=303)

    @app.post("/fortify/create")
    async def fortify_create(request: Request):
        import json as _json

        from paladin.fortify.campaign import SSCCampaignError, create_campaign_from_group
        from paladin.fortify.discovery import Group, SubApp

        form = await request.form()
        if form.get("token") != ui_token:
            raise HTTPException(403, "Jeton d'interface invalide : recharger la page.")
        keys = [str(k) for k in form.getlist("group")]
        cached = _json.loads(_discovery_file().read_text(encoding="utf-8"))
        repos = [Path(x.strip()) for x in str(form.get("repos", "")).splitlines() if x.strip()]
        workbook = str(form.get("workbook", "")).strip()
        created, errors = [], []
        for g in cached["groups"]:
            if g["key"] not in keys:
                continue
            group = Group(g["key"], [SubApp(**a) for a in g["apps"]])
            try:
                created.append(
                    create_campaign_from_group(
                        settings, conn, group, repo_paths=repos, target_workbook=Path(workbook) if workbook else None
                    )
                )
            except SSCCampaignError as exc:
                errors.append(f"{g['key']} : {exc}")
        if errors:
            _gflash("Certaines entrées n'ont pas été créées.", "warn", errors)
        if len(created) == 1 and not errors:
            _flash(
                conn,
                created[0],
                "Campagne créée à partir de SSC. Étape suivante : « Importer / réimporter les sources ».",
                "ok",
            )
            return RedirectResponse(f"/c/{created[0]}", status_code=303)
        if created:
            _gflash(f"Campagnes créées : {', '.join(created)}. Ouvrir chacune et cliquer « Importer ».", "ok")
        return RedirectResponse("/fortify/discover", status_code=303)

    @app.get("/health")
    def health() -> dict:
        return {"status": "ok", "version": __version__}

    @app.get("/", response_class=HTMLResponse)
    def index(request: Request):
        campaigns = store.list_campaigns(conn)
        if len(campaigns) == 1:
            return RedirectResponse(f"/c/{campaigns[0]['id']}", status_code=303)
        cards = [{**c, "counters": q.counters(conn, c["id"])} for c in campaigns]
        return render(request, "index.html", campaigns=cards, home=str(settings.home))

    @app.get("/c/{cid}", response_class=HTMLResponse)
    def dashboard(request: Request, cid: str):
        from paladin.importers import profiles

        campaign = _campaign(cid)
        tools = []
        for t in store.list_tools(conn, cid):
            run = conn.execute(
                "SELECT * FROM import_run WHERE tool_id = ? ORDER BY started_at DESC LIMIT 1", (t["id"],)
            ).fetchone()
            count = conn.execute("SELECT COUNT(*) FROM finding WHERE tool_id = ?", (t["id"],)).fetchone()[0]
            tools.append(
                {
                    **t,
                    "count": count,
                    "run": dict(run) if run else None,
                    "report": loads(run["report_json"], {}) if run else {},
                }
            )
        last_export = conn.execute(
            "SELECT * FROM export_run WHERE campaign_id = ? ORDER BY created_at DESC LIMIT 1", (cid,)
        ).fetchone()
        last = store.get_ui_state(conn, cid, "last_finding")
        return render(
            request,
            "campaign.html",
            campaign=campaign,
            counters=q.counters(conn, cid),
            tools=tools,
            pending_profiles=[p for p in profiles.list_profiles(conn, cid) if p["status"] == "proposed"],
            last_export=dict(last_export) if last_export else None,
            last_export_summary=loads(last_export["summary_json"], {}) if last_export else {},
            last_finding=last,
            last_decided=q.last_decided(conn, cid),
            flash=_pop_flash(conn, cid),
            home=str(settings.home),
            agent=agent_jobs.status(conn, cid),
        )

    @app.get("/c/{cid}/queue", response_class=HTMLResponse)
    def queue_page(request: Request, cid: str, view: str = "todo", tool: str = "", search: str = ""):
        campaign = _campaign(cid)
        rows = q.list_findings(conn, cid, view, tool or None, search or None)
        return render(
            request,
            "queue.html",
            campaign=campaign,
            rows=rows,
            view=view,
            tool=tool,
            search=search,
            tools=[t["label"] for t in store.list_tools(conn, cid)],
            counters=q.counters(conn, cid),
            flash=_pop_flash(conn, cid),
        )

    @app.get("/c/{cid}/next", response_class=HTMLResponse)
    def start_view(cid: str, view: str = "todo", tool: str = ""):
        rows = q.list_findings(conn, cid, view, tool or None, limit=1)
        if not rows:
            _flash(conn, cid, f"Rien dans la vue « {q.VIEWS.get(view, view)} ».")
            return RedirectResponse(f"/c/{cid}/queue?{urlencode({'view': view, 'tool': tool})}", status_code=303)
        return RedirectResponse(
            f"/c/{cid}/f/{rows[0]['id']}?{urlencode({'view': view, 'tool': tool})}", status_code=303
        )

    @app.get("/c/{cid}/f/{fid}", response_class=HTMLResponse)
    def card_page(request: Request, cid: str, fid: str, view: str = "todo", tool: str = ""):
        campaign = _campaign(cid)
        try:
            c = q.card(conn, cid, fid, view, tool or None)
        except LookupError:
            raise HTTPException(404, "Finding inconnu") from None
        store.set_ui_state(
            conn, cid, "last_finding", {"id": fid, "source_id": c.finding["source_id"], "view": view, "tool": tool}
        )
        prev_id = q.neighbour(conn, cid, fid, view, tool or None, -1)
        next_id = q.neighbour(conn, cid, fid, view, tool or None, +1)
        return render(
            request,
            "card.html",
            campaign=campaign,
            c=c,
            view=view,
            tool=tool,
            prev_id=prev_id,
            next_id=next_id,
            counters=q.counters(conn, cid),
            flash=_pop_flash(conn, cid),
            editor=bool(settings.editor.get("open_command")),
        )

    # ---------------------------------------------------------------- actions

    def _after(cid: str, fid: str, view: str, tool: str, move: bool) -> RedirectResponse:
        params = urlencode({"view": view, "tool": tool})
        if move:
            nxt = q.neighbour(conn, cid, fid, view, tool or None, +1)
            if nxt:
                return RedirectResponse(f"/c/{cid}/f/{nxt}?{params}", status_code=303)
            _flash(conn, cid, f"Fin de la vue « {q.VIEWS.get(view, view)} ».", "ok")
            return RedirectResponse(f"/c/{cid}/queue?{params}", status_code=303)
        return RedirectResponse(f"/c/{cid}/f/{fid}?{params}", status_code=303)

    @app.post("/c/{cid}/f/{fid}/decide")
    async def decide(request: Request, cid: str, fid: str):
        from paladin.analysis import latest_analysis

        _campaign(cid)
        form = await _form(request)
        view, tool, op = form.get("view", "todo"), form.get("tool", ""), form.get("op", "")
        try:
            revision = int(form.get("revision", "0"))
            if op == "skip":
                dec.record_decision(conn, fid, expected_revision=revision, action=DecisionAction.SKIP, author=AUTHOR)
                return _after(cid, fid, view, tool, move=True)
            if op == "investigate":
                dec.record_decision(
                    conn,
                    fid,
                    expected_revision=revision,
                    action=DecisionAction.INVESTIGATE,
                    author=AUTHOR,
                    investigation_question=form.get("question"),
                    investigation_reason=form.get("reason"),
                    analysis_id=form.get("analysis_id") or None,
                )
                _flash(conn, cid, "Mis « À investiguer » — décision enregistrée, aucune valeur Excel.", "ok")
                return _after(cid, fid, view, tool, move=True)
            if op == "decide":
                verdict = form.get("verdict") or None
                if verdict not in ("TRUE_POSITIVE", "NOT_AN_ISSUE"):
                    raise dec.DecisionError("Choisir True Positive ou Not an issue (ou « À investiguer »).")
                comment = form.get("comment", "")
                analysis = latest_analysis(conn, fid)
                accepted = (
                    analysis is not None
                    and analysis["proposed_verdict"] == verdict
                    and comment.strip() in ("", (analysis["suggested_comment"] or "").strip())
                )
                discussion = form.get("discussion") == "1" or comment.strip() == DISCUSSION_COMMENT
                dec.record_decision(
                    conn,
                    fid,
                    expected_revision=revision,
                    action=DecisionAction.ACCEPT if accepted else DecisionAction.CORRECT,
                    author=AUTHOR,
                    verdict=Verdict(verdict),
                    comment=comment,
                    analysis_id=analysis["id"] if analysis else None,
                    discussion_required=discussion,
                    discussion_reason=form.get("discussion_reason") or None,
                    correction_category=None if accepted else (form.get("correction_category") or None),
                )
                f = conn.execute("SELECT export_state FROM finding WHERE id = ?", (fid,)).fetchone()
                _flash(
                    conn,
                    cid,
                    f"Décision enregistrée ({VERDICT_LABELS[verdict]}). {EXPORT_LABELS[f['export_state']]}.",
                    "ok",
                )
                return _after(cid, fid, view, tool, move=form.get("next") != "0")
            raise dec.DecisionError("Action inconnue.")
        except ConflictError as exc:
            _flash(conn, cid, f"{exc}", "warn")
        except dec.DecisionError as exc:
            _flash(conn, cid, str(exc), "warn")
        return _after(cid, fid, view, tool, move=False)

    @app.post("/c/{cid}/undo")
    async def undo(request: Request, cid: str):
        form = await _form(request)
        last = q.last_decided(conn, cid)
        if last is None:
            _flash(conn, cid, "Aucune décision à annuler.", "warn")
            return RedirectResponse(f"/c/{cid}", status_code=303)
        f = conn.execute("SELECT revision FROM finding WHERE id = ?", (last["finding_id"],)).fetchone()
        dec.undo_last(conn, last["finding_id"], expected_revision=f["revision"], author=AUTHOR)
        state = conn.execute("SELECT export_state FROM finding WHERE id = ?", (last["finding_id"],)).fetchone()[0]
        extra = " L'Excel déjà exporté est PÉRIMÉ : réexporter." if state == "stale" else ""
        _flash(conn, cid, f"Décision annulée sur {last['source_id']} (historique conservé).{extra}", "ok")
        params = urlencode({"view": form.get("view", "all"), "tool": form.get("tool", "")})
        return RedirectResponse(f"/c/{cid}/f/{last['finding_id']}?{params}", status_code=303)

    @app.post("/api/c/{cid}/f/{fid}/draft")
    async def draft(request: Request, cid: str, fid: str):
        if request.headers.get("x-paladin-token") != ui_token:
            return JSONResponse({"error": "jeton"}, status_code=403)
        body = await request.json()
        row = conn.execute("SELECT revision FROM finding WHERE id = ? AND campaign_id = ?", (fid, cid)).fetchone()
        if row is None:
            return JSONResponse({"error": "finding"}, status_code=404)
        if str(body.get("revision")) != str(row["revision"]):
            return JSONResponse({"error": "révision périmée"}, status_code=409)  # brouillon arrivé après une décision
        verdict = body.get("verdict") if body.get("verdict") in ("TRUE_POSITIVE", "NOT_AN_ISSUE") else None
        dec.save_draft(conn, fid, verdict, (body.get("comment") or "")[:4000])
        return JSONResponse({"saved": True})

    @app.post("/c/{cid}/agent/enqueue")
    async def agent_enqueue(request: Request, cid: str):
        await _form(request)
        n = agent_jobs.enqueue_analysis(conn, cid)
        _flash(conn, cid, f"{n} finding(s) mis en file pour l'agent." if n else "Aucun finding à mettre en file.", "ok")
        return RedirectResponse(f"/c/{cid}", status_code=303)

    @app.get("/api/c/{cid}/f/{fid}/status")
    def finding_status(cid: str, fid: str):
        row = conn.execute(
            "SELECT f.revision, f.processing_state, (SELECT MAX(seq) FROM analysis a WHERE a.finding_id = f.id) AS seq"
            " FROM finding f WHERE f.id = ? AND f.campaign_id = ?",
            (fid, cid),
        ).fetchone()
        if row is None:
            raise HTTPException(404, "Finding inconnu")
        return {
            "revision": row["revision"],
            "processing_state": row["processing_state"],
            "analysis_seq": row["seq"] or 0,
        }

    # ------------------------------------------------------- rapprochement inter-outils

    def _missing_comparative(campaign: dict) -> dict[str, list[str]]:
        """Colonnes `Found in` / `criticality in` absentes des onglets existants."""
        from openpyxl import load_workbook

        from paladin.contracts import criticality_in_header, found_in_header
        from paladin.excel.export import target_path

        tools = store.list_tools(conn, campaign["id"])
        path = target_path(settings, campaign)
        if not path.exists():
            return {}
        wb = load_workbook(path, read_only=True)
        try:
            out = {}
            ext = campaign["config"].get("sheet_extensions") or {}
            for t in tools:
                if not t["sheet_name"] or t["sheet_name"] not in wb.sheetnames:
                    continue
                headers = {
                    str(c.value).strip().lower() for c in next(wb[t["sheet_name"]].iter_rows(max_row=1)) if c.value
                }
                headers |= {h.lower() for h in ext.get(t["sheet_name"], [])}
                missing = [
                    h
                    for o in tools
                    if o["id"] != t["id"]
                    for h in (found_in_header(o["label"]), criticality_in_header(o["label"]))
                    if h.lower() not in headers
                ]
                if missing:
                    out[t["sheet_name"]] = missing
            return out
        finally:
            wb.close()

    @app.get("/c/{cid}/match", response_class=HTMLResponse)
    def match_page(request: Request, cid: str):
        from paladin import matching

        campaign = _campaign(cid)
        rels = conn.execute(
            "SELECT r.*, fa.source_id AS a_sid, fb.source_id AS b_sid, ta.label AS a_tool, tb.label AS b_tool,"
            " fa.normalized_path AS a_path, fa.line_number AS a_line, fb.line_number AS b_line,"
            " (SELECT action FROM relation_event e WHERE e.relation_id = r.id ORDER BY e.created_at DESC LIMIT 1)"
            " AS last_action FROM finding_relation r JOIN finding fa ON fa.id = r.finding_a_id"
            " JOIN finding fb ON fb.id = r.finding_b_id JOIN tool ta ON ta.id = fa.tool_id"
            " JOIN tool tb ON tb.id = fb.tool_id WHERE fa.campaign_id = ?"
            " ORDER BY CASE r.state WHEN 'reexam_required' THEN 0 WHEN 'proposed' THEN 1 ELSE 2 END,"
            " (SELECT action FROM relation_event e WHERE e.relation_id = r.id ORDER BY e.created_at DESC LIMIT 1)"
            " = 'defer', r.score DESC, fa.source_id",
            (cid,),
        ).fetchall()
        runs = conn.execute(
            "SELECT c.*, ta.label AS a, tb.label AS b FROM comparison_run c JOIN tool ta ON ta.id = c.tool_a_id"
            " JOIN tool tb ON tb.id = c.tool_b_id WHERE c.campaign_id = ? AND c.created_at = (SELECT MAX(created_at)"
            " FROM comparison_run c2 WHERE c2.tool_a_id = c.tool_a_id AND c2.tool_b_id = c.tool_b_id)",
            (cid,),
        ).fetchall()
        return render(
            request,
            "match.html",
            campaign=campaign,
            rels=[dict(r) for r in rels],
            runs=[dict(r) | {"corpus": loads(r["corpus_json"], {})} for r in runs],
            exact=matching.exact_candidates(conn, cid),
            mcount=matching.counters(conn, cid),
            missing=_missing_comparative(campaign),
            counters=q.counters(conn, cid),
            flash=_pop_flash(conn, cid),
        )

    @app.post("/c/{cid}/match/run")
    async def match_run(request: Request, cid: str):
        from paladin import matching

        await _form(request)
        results = matching.compare_all(conn, cid)
        lines = [
            f"{r.tool_a} ↔ {r.tool_b} : {r.new} nouveau(x) candidat(s), {r.unchanged} inchangé(s),"
            f" {r.reexam} à réexaminer, corpus {r.completeness}" + "".join(f" — {n}" for n in r.notes)
            for r in results
        ]
        _flash(conn, cid, "Rapprochement préparé." if results else "Il faut au moins deux outils.", "ok", lines)
        return RedirectResponse(f"/c/{cid}/match", status_code=303)

    @app.post("/c/{cid}/match/batch")
    async def match_batch(request: Request, cid: str):
        from paladin import matching

        form = await request.form()
        if form.get("token") != ui_token:
            raise HTTPException(403, "Jeton d'interface invalide : recharger la page.")
        ids = [str(x) for x in form.getlist("rid")]
        try:
            batch = matching.confirm_batch(conn, cid, ids, AUTHOR)
            _flash(conn, cid, f"{len(ids)} lien(s) confirmé(s) en lot (lot {batch[:8]}), un événement par lien.", "ok")
        except matching.MatchingError as exc:
            _flash(conn, cid, str(exc), "warn")
        return RedirectResponse(f"/c/{cid}/match", status_code=303)

    @app.post("/c/{cid}/match/extensions")
    async def match_extensions(request: Request, cid: str):
        await _form(request)
        campaign = _campaign(cid)
        cfg = campaign["config"]
        ext = cfg.setdefault("sheet_extensions", {})
        for sheet, headers in _missing_comparative(campaign).items():
            ext[sheet] = [*ext.get(sheet, []), *headers]
        store.update_campaign_config(conn, cid, cfg)
        _flash(
            conn,
            cid,
            "Colonnes comparatives ajoutées à l'export (en fin de ligne d'en-tête, rien n'est déplacé).",
            "ok",
        )
        return RedirectResponse(f"/c/{cid}/match", status_code=303)

    @app.get("/c/{cid}/match/{rid}", response_class=HTMLResponse)
    def match_card(request: Request, cid: str, rid: str):
        from paladin.analysis import excerpt_for_finding
        from paladin.review.decisions import current_decision

        campaign = _campaign(cid)
        rel = conn.execute("SELECT * FROM finding_relation WHERE id = ?", (rid,)).fetchone()
        if rel is None:
            raise HTTPException(404, "Lien inconnu")
        sides = []
        for fid in (rel["finding_a_id"], rel["finding_b_id"]):
            f = conn.execute(
                "SELECT f.*, t.label AS tool_label FROM finding f JOIN tool t ON t.id = f.tool_id WHERE f.id = ?",
                (fid,),
            ).fetchone()
            sides.append(
                {
                    "f": dict(f),
                    "cwe": loads(f["cwe_ids_json"], []),
                    "decision": current_decision(conn, fid),
                    "excerpt": excerpt_for_finding(conn, f),
                }
            )
        history = conn.execute(
            "SELECT * FROM relation_event WHERE relation_id = ? ORDER BY created_at DESC", (rid,)
        ).fetchall()
        return render(
            request,
            "match_card.html",
            campaign=campaign,
            rel=dict(rel),
            sides=sides,
            evidence=loads(rel["evidence_json"], []),
            differences=loads(rel["differences_json"], []),
            history=[dict(h) for h in history],
            counters=q.counters(conn, cid),
            flash=_pop_flash(conn, cid),
        )

    @app.post("/c/{cid}/match/{rid}/decide")
    async def match_decide(request: Request, cid: str, rid: str):
        from paladin import matching

        form = await _form(request)
        op = form.get("op", "")
        mapping = {
            "same": ("confirm", "same_occurrence"),
            "root": ("confirm", "same_root_cause"),
            "different": ("reject", None),
            "defer": ("defer", None),
            "undo": ("undo", None),
        }
        if op not in mapping:
            raise HTTPException(400, "Action inconnue")
        action, kind = mapping[op]
        try:
            matching.decide(
                conn, rid, action, author=AUTHOR, expected_revision=int(form.get("revision", "0")), relation_type=kind
            )
        except (matching.MatchingError, ConflictError) as exc:
            _flash(conn, cid, str(exc), "warn")
            return RedirectResponse(f"/c/{cid}/match/{rid}", status_code=303)
        if op == "undo":
            _flash(conn, cid, "Décision de rapprochement annulée (historique conservé).", "ok")
            return RedirectResponse(f"/c/{cid}/match/{rid}", status_code=303)
        nxt = conn.execute(
            "SELECT r.id FROM finding_relation r JOIN finding f ON f.id = r.finding_a_id WHERE f.campaign_id = ?"
            " AND r.state IN ('proposed', 'reexam_required') AND r.id != ? AND COALESCE((SELECT action FROM"
            " relation_event e WHERE e.relation_id = r.id ORDER BY e.created_at DESC LIMIT 1), '') != 'defer'"
            " ORDER BY r.state = 'proposed', r.score DESC LIMIT 1",
            (cid, rid),
        ).fetchone()
        _flash(conn, cid, "Lien enregistré. Aucun verdict n'est propagé entre outils.", "ok")
        return RedirectResponse(f"/c/{cid}/match/{nxt['id']}" if nxt else f"/c/{cid}/match", status_code=303)

    # --------------------------------------------------------- nouvel onglet (schéma)

    def _sheet_names(campaign: dict) -> list[str]:
        from openpyxl import load_workbook

        from paladin.excel.export import target_path

        path = target_path(settings, campaign)
        if not path.exists():
            return []
        wb = load_workbook(path, read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()

    @app.get("/c/{cid}/tools/{label}/schema", response_class=HTMLResponse)
    def schema_page(request: Request, cid: str, label: str):
        from paladin.excel import sheets

        campaign = _campaign(cid)
        try:
            proposal = sheets.propose(conn, cid, label, _sheet_names(campaign))
        except (sheets.SheetSchemaError, store.NotFoundError) as exc:
            _flash(conn, cid, str(exc), "warn")
            return RedirectResponse(f"/c/{cid}", status_code=303)
        return render(
            request,
            "schema.html",
            campaign=campaign,
            proposal=proposal,
            label=label,
            preview=sheets.preview(conn, cid, proposal),
            counters=q.counters(conn, cid),
            flash=_pop_flash(conn, cid),
        )

    @app.post("/c/{cid}/tools/{label}/schema")
    async def schema_validate(request: Request, cid: str, label: str):
        from paladin.excel import sheets

        form = await _form(request)
        campaign = _campaign(cid)
        existing = _sheet_names(campaign)
        try:
            proposal = sheets.apply_edits(sheets.propose(conn, cid, label, existing), form)
            sheets.validate(proposal, existing)
            version = sheets.save_validated(conn, cid, proposal, "human")
        except sheets.SheetSchemaError as exc:
            _flash(conn, cid, f"Schéma non validé : {exc}", "warn")
            return RedirectResponse(f"/c/{cid}/tools/{label}/schema", status_code=303)
        _flash(
            conn, cid, f"Onglet « {proposal.sheet_name} » validé (v{version}). Il sera créé au prochain export.", "ok"
        )
        return RedirectResponse(f"/c/{cid}", status_code=303)

    # ------------------------------------------------- mémoire, règles, groupes, mesures

    @app.post("/c/{cid}/f/{fid}/reference")
    async def toggle_reference(request: Request, cid: str, fid: str):
        from paladin import calibration

        form = await _form(request)
        value = form.get("reference") == "1"
        back = form.get("back") or f"/c/{cid}/f/{fid}?view={form.get('view', 'all')}"
        if not back.startswith(f"/c/{cid}/"):
            back = f"/c/{cid}"
        if conn.execute("SELECT 1 FROM finding WHERE id = ? AND campaign_id = ?", (fid, cid)).fetchone() is None:
            raise HTTPException(404, "Finding inconnu")
        try:
            calibration.set_role(conn, fid, value)
        except calibration.CalibrationError as exc:
            _flash(conn, cid, str(exc), "warn")
            return RedirectResponse(back, status_code=303)
        _flash(
            conn,
            cid,
            "Ajouté au jeu de référence : sa décision ne sera jamais montrée à l'agent."
            if value
            else "Retiré du jeu de référence : il devient un exemple, montrable à l'agent.",
            "ok",
        )
        return RedirectResponse(back, status_code=303)

    @app.get("/c/{cid}/rules", response_class=HTMLResponse)
    def rules_page(request: Request, cid: str, finding: str = ""):
        from paladin import rules as rules_mod

        campaign = _campaign(cid)
        draft = None
        if finding:
            f = conn.execute(
                "SELECT f.*, t.label AS tool_label FROM finding f JOIN tool t ON t.id = f.tool_id"
                " WHERE f.id = ? AND f.campaign_id = ?",
                (finding, cid),
            ).fetchone()
            if f is not None:
                dec = None
                if f["current_decision_id"]:
                    dec = conn.execute(
                        "SELECT * FROM decision_event WHERE id = ?", (f["current_decision_id"],)
                    ).fetchone()
                draft = {
                    "finding": dict(f),
                    "decision": dict(dec) if dec else None,
                    "conditions": rules_mod.conditions_from_finding(f, f["tool_label"]),
                }
        all_rules = rules_mod.list_rules(conn, cid)
        derived = {
            r.id: conn.execute(
                "SELECT COUNT(*) FROM finding f JOIN decision_event e ON e.id = f.current_decision_id"
                " WHERE e.rule_id = ?",
                (r.id,),
            ).fetchone()[0]
            for r in all_rules
        }
        return render(
            request,
            "rules.html",
            campaign=campaign,
            rules=all_rules,
            draft=draft,
            derived=derived,
            counters=q.counters(conn, cid),
            flash=_pop_flash(conn, cid),
        )

    @app.post("/c/{cid}/rules/create")
    async def rule_create(request: Request, cid: str):
        from paladin import rules as rules_mod

        form = await _form(request)
        conditions = {k: form.get(f"cond_{k}", "").strip() for k in rules_mod.CONDITION_KEYS}
        exceptions = [x.strip() for x in form.get("exceptions", "").split(",") if x.strip()]
        try:
            rule = rules_mod.create_from_decision(
                conn,
                form.get("finding_id", ""),
                title=form.get("title", ""),
                author=AUTHOR,
                conditions={k: v for k, v in conditions.items() if v},
                exceptions=exceptions,
                counter_example=form.get("counter_example") or None,
                application_scope=form.get("app_scope") == "1",
            )
        except rules_mod.RuleError as exc:
            _flash(conn, cid, str(exc), "warn")
            return RedirectResponse(f"/c/{cid}/rules?finding={form.get('finding_id', '')}", status_code=303)
        _flash(conn, cid, f"Règle « {rule.title} » proposée. Elle ne s'appliquera qu'après validation.", "ok")
        return RedirectResponse(f"/c/{cid}/rules", status_code=303)

    @app.post("/c/{cid}/rules/{rid}/{op}")
    async def rule_op(request: Request, cid: str, rid: str, op: str):
        from paladin import rules as rules_mod

        form = await _form(request)
        try:
            if op == "validate":
                rule = rules_mod.validate(conn, rid, AUTHOR)
                _flash(
                    conn,
                    cid,
                    f"Règle « {rule.title} » active : elle apparaît sur les fiches concernées et peut"
                    " servir de base à un lot.",
                    "ok",
                )
            elif op == "revoke":
                n = rules_mod.revoke(conn, rid, AUTHOR, form.get("reason") or "révoquée par l'analyste")
                _flash(
                    conn,
                    cid,
                    f"Règle révoquée : {n} décision(s) dérivée(s) passée(s) en « réexamen requis »"
                    " (historique conservé).",
                    "ok",
                )
            else:
                raise HTTPException(400, "Action inconnue")
        except rules_mod.RuleError as exc:
            _flash(conn, cid, str(exc), "warn")
        return RedirectResponse(f"/c/{cid}/rules", status_code=303)

    @app.get("/c/{cid}/groups", response_class=HTMLResponse)
    def groups_page(request: Request, cid: str, key: str = ""):
        from paladin import groups as groups_mod

        campaign = _campaign(cid)
        all_groups = groups_mod.groups(conn, cid)
        selected = next((g for g in all_groups if g.key == key), None)
        return render(
            request,
            "groups.html",
            campaign=campaign,
            groups=all_groups,
            selected=selected,
            history=groups_mod.history(conn, cid),
            counters=q.counters(conn, cid),
            flash=_pop_flash(conn, cid),
        )

    @app.post("/c/{cid}/groups/batch")
    async def groups_batch(request: Request, cid: str):
        from paladin import groups as groups_mod

        form = await request.form()
        if form.get("token") != ui_token:
            raise HTTPException(403, "Jeton d'interface invalide : recharger la page.")
        key = str(form.get("key", ""))
        frozen = {str(fid): int(str(form.get(f"rev_{fid}", "-1"))) for fid in form.getlist("member")}
        try:
            batch_id = groups_mod.execute(
                conn,
                cid,
                key,
                frozen,
                verdict=str(form.get("verdict", "")),
                comment=str(form.get("comment", "")).strip() or None,
                author=AUTHOR,
            )
        except groups_mod.BatchError as exc:
            _flash(conn, cid, str(exc), "warn")
            return RedirectResponse(f"/c/{cid}/groups?{urlencode({'key': key})}", status_code=303)
        _flash(
            conn,
            cid,
            f"Lot {batch_id[:8]} : {len(frozen)} décision(s) enregistrée(s), une par membre, marquées"
            " « lot ». Annulable depuis l'historique des lots.",
            "ok",
        )
        return RedirectResponse(f"/c/{cid}/groups", status_code=303)

    @app.post("/c/{cid}/batches/{bid}/undo")
    async def batch_undo(request: Request, cid: str, bid: str):
        from paladin import groups as groups_mod

        await _form(request)
        try:
            report = groups_mod.undo(conn, bid, AUTHOR)
            details = [f"{k} : {', '.join(v) or '—'}" for k, v in report.items()]
            _flash(
                conn, cid, "Lot annulé (historique conservé). Si l'Excel était exporté, il est périmé.", "ok", details
            )
        except groups_mod.BatchError as exc:
            _flash(conn, cid, str(exc), "warn")
        return RedirectResponse(f"/c/{cid}/groups", status_code=303)

    @app.get("/c/{cid}/stats", response_class=HTMLResponse)
    def stats_page(request: Request, cid: str, search: str = ""):
        from paladin.review import memory

        campaign = _campaign(cid)
        return render(
            request,
            "stats.html",
            campaign=campaign,
            m=memory.pilot_metrics(conn, cid),
            search=search,
            results=memory.search(conn, cid, search) if search.strip() else None,
            counters=q.counters(conn, cid),
            flash=_pop_flash(conn, cid),
        )

    # ------------------------------------------------- calibration sur analyses manuelles

    def _manual_args(params: Any) -> dict[str, Any]:
        """Lecture choisie dans l'aperçu : colonnes (`col_<rôle>`) et traductions (`vk_<n>` / `vv_<n>`)."""
        from paladin.excel import manual

        mapping = None
        if any(k.startswith("col_") for k in params):
            mapping = {
                r: int(params[f"col_{r}"]) if params.get(f"col_{r}", "") not in ("", "-1") else None
                for r in manual.ROLES
            }
        value_map = {
            params[k]: params.get("vv_" + k[3:], manual.IGNORE) for k in params if k.startswith("vk_") and params[k]
        }
        return {
            "sheet": params.get("sheet") or None,
            "tool_label": params.get("tool") or None,
            "mapping": mapping,
            "value_map": value_map or None,
        }

    @app.get("/c/{cid}/calibration", response_class=HTMLResponse)
    def calibration_page(request: Request, cid: str, version: int | None = None):
        from paladin import calibration

        campaign = _campaign(cid)
        return render(
            request,
            "calibration.html",
            campaign=campaign,
            r=calibration.report(conn, cid, version),
            gain=calibration.time_gain(conn, cid),
            conventions=calibration.current_conventions(conn),
            history=calibration.conventions_history(conn),
            examples=calibration.examples_summary(conn),
            agent=agent_jobs.status(conn, cid),
            counters=q.counters(conn, cid),
            flash=_pop_flash(conn, cid),
        )

    @app.post("/c/{cid}/calibration/upload")
    async def calibration_upload(request: Request, cid: str):
        from paladin.excel import manual

        _campaign(cid)
        form = await request.form()
        if form.get("token") != ui_token:
            raise HTTPException(403, "Jeton d'interface invalide : recharger la page.")
        upload = form.get("file")
        if upload is None or isinstance(upload, str) or not upload.filename:
            _flash(conn, cid, "Choisir un classeur .xlsx à déposer.", "warn")
            return RedirectResponse(f"/c/{cid}/calibration", status_code=303)
        path = manual.save_upload(settings.campaign_dir(cid), upload.filename, await upload.read())
        return RedirectResponse(f"/c/{cid}/calibration/import?file={path.name}", status_code=303)

    @app.get("/c/{cid}/calibration/import", response_class=HTMLResponse)
    def calibration_preview(request: Request, cid: str, file: str):
        from paladin.excel import manual

        campaign = _campaign(cid)
        args = _manual_args(dict(request.query_params))
        try:
            path = manual.stored_copy(settings.campaign_dir(cid), file)
            plan = manual.scan(conn, cid, path, **args)
        except manual.ManualImportError as exc:
            _flash(conn, cid, str(exc), "warn")
            return RedirectResponse(f"/c/{cid}/calibration", status_code=303)
        statuses: dict[str, list] = {}
        for it in plan.items:
            statuses.setdefault(it.status, []).append(it)
        return render(
            request,
            "manual_import.html",
            campaign=campaign,
            plan=plan,
            file=file,
            tool=args["tool_label"] or "",
            tools=store.list_tools(conn, cid),
            statuses=statuses,
            ROLES=manual.ROLES,
            VALUE_CHOICES=manual.VALUE_CHOICES,
            counters=q.counters(conn, cid),
            flash=_pop_flash(conn, cid),
        )

    @app.post("/c/{cid}/calibration/import")
    async def calibration_import(request: Request, cid: str):
        from paladin.excel import manual

        form = await _form(request)
        try:
            path = manual.stored_copy(settings.campaign_dir(cid), form.get("file", ""))
            plan = manual.scan(conn, cid, path, **_manual_args(form))
            res = manual.apply(conn, cid, plan, AUTHOR, role=form.get("role", "auto"))
        except manual.ManualImportError as exc:
            _flash(conn, cid, str(exc), "warn")
            return RedirectResponse(f"/c/{cid}/calibration", status_code=303)
        details = [f"{k} : {v}" for k, v in res.roles.items() if v]
        if res.baseline:
            details.append(
                f"Temps manuel enregistré : {res.baseline['minutes']} min pour {res.baseline['findings']} finding(s)"
            )
        conflicts = plan.by_status("conflit")
        if conflicts:
            details.append(f"{len(conflicts)} conflit(s) : la décision Paladin est conservée")
        _flash(
            conn,
            cid,
            f"{len(res.imported)} analyse(s) manuelle(s) importée(s) (lot {res.batch_id[:8]}, annulable depuis"
            " « Groupes et lots »). Votre classeur n'a pas été modifié.",
            "ok",
            details,
        )
        return RedirectResponse(f"/c/{cid}/calibration", status_code=303)

    @app.post("/c/{cid}/calibration/resplit")
    async def calibration_resplit(request: Request, cid: str):
        from paladin import calibration

        await _form(request)
        roles = calibration.resplit(conn, cid)
        _flash(conn, cid, "Nouvelle répartition exemples / référence.", "ok", [f"{k} : {v}" for k, v in roles.items()])
        return RedirectResponse(f"/c/{cid}/calibration", status_code=303)

    @app.post("/c/{cid}/calibration/conventions")
    async def calibration_conventions(request: Request, cid: str):
        from paladin import calibration

        form = await _form(request)
        try:
            v = calibration.save_conventions(conn, form.get("text", ""), AUTHOR, form.get("note", ""))
        except calibration.CalibrationError as exc:
            _flash(conn, cid, str(exc), "warn")
        else:
            _flash(conn, cid, f"Conventions v{v} enregistrées : l'agent les reçoit dès sa prochaine analyse.", "ok")
        return RedirectResponse(f"/c/{cid}/calibration#conventions", status_code=303)

    @app.post("/c/{cid}/calibration/blind")
    async def calibration_blind(request: Request, cid: str):
        from paladin import calibration

        await _form(request)
        n = calibration.enqueue_blind(conn, cid)
        _flash(
            conn,
            cid,
            f"{n} analyse(s) à l'aveugle en file. Lancer l'agent (page « Agent » ou `Paladin.cmd agent run`)."
            if n
            else "Rien à mettre en file : toute la référence est déjà analysée avec ces conventions.",
            "ok" if n else "info",
        )
        return RedirectResponse(f"/c/{cid}/calibration", status_code=303)

    @app.post("/c/{cid}/calibration/baseline")
    async def calibration_baseline(request: Request, cid: str):
        from paladin import calibration

        form = await _form(request)
        try:
            calibration.set_baseline(
                conn, cid, float(form.get("minutes", "0").replace(",", ".")), int(form.get("findings", "0"))
            )
        except (ValueError, calibration.CalibrationError):
            _flash(conn, cid, "Indiquer un temps (minutes) et un nombre de findings positifs.", "warn")
        else:
            _flash(conn, cid, "Temps manuel de référence enregistré.", "ok")
        return RedirectResponse(f"/c/{cid}/calibration#temps", status_code=303)

    # ------------------------------------------------- classeur cible et reprise

    @app.post("/c/{cid}/excel/target")
    async def excel_target(request: Request, cid: str):
        from paladin.excel import reprise

        form = await _form(request)
        try:
            missing = reprise.change_target(settings, conn, cid, Path(form.get("path", "").strip()))
        except reprise.RepriseError as exc:
            _flash(conn, cid, str(exc), "warn")
            return RedirectResponse(f"/c/{cid}", status_code=303)
        details = [f"Onglets attendus absents : {', '.join(missing)}"] if missing else []
        details.append("Si ce classeur contient déjà des verdicts : « Reprendre les décisions déjà saisies ».")
        _flash(conn, cid, "Classeur cible changé.", "warn" if missing else "ok", details)
        return RedirectResponse(f"/c/{cid}", status_code=303)

    @app.get("/c/{cid}/reprise", response_class=HTMLResponse)
    def reprise_page(request: Request, cid: str):
        from paladin.excel import reprise

        campaign = _campaign(cid)
        try:
            plan = reprise.scan(settings, conn, cid)
        except reprise.RepriseError as exc:
            _flash(conn, cid, str(exc), "warn")
            return RedirectResponse(f"/c/{cid}", status_code=303)
        statuses = {}
        for it in plan.items:
            statuses.setdefault(it.status, []).append(it)
        return render(
            request,
            "reprise.html",
            campaign=campaign,
            plan=plan,
            statuses=statuses,
            counters=q.counters(conn, cid),
            flash=_pop_flash(conn, cid),
        )

    @app.post("/c/{cid}/reprise")
    async def reprise_apply(request: Request, cid: str):
        from paladin.excel import reprise

        await _form(request)
        try:
            rid, plan = reprise.apply(settings, conn, cid, AUTHOR)
        except reprise.RepriseError as exc:
            _flash(conn, cid, str(exc), "warn")
            return RedirectResponse(f"/c/{cid}/reprise", status_code=303)
        conflicts = plan.by_status("conflit")
        _flash(
            conn,
            cid,
            f"{len(plan.importable)} décision(s) reprise(s) du classeur (reprise {rid[:8]}),"
            " marquées « reprise Excel ». Annulable en bloc depuis « Groupes et lots ».",
            "ok",
            [f"{len(conflicts)} conflit(s) conservé(s) côté Paladin"] if conflicts else None,
        )
        return RedirectResponse(f"/c/{cid}", status_code=303)

    @app.post("/c/{cid}/import")
    async def run_import(request: Request, cid: str):
        from paladin.importers.pipeline import import_tool

        form = await _form(request)
        labels = [form["tool"]] if form.get("tool") else [t["label"] for t in store.list_tools(conn, cid)]
        lines, level = [], "ok"
        for label in labels:
            report = import_tool(settings, conn, cid, label)
            line = report.summary() + "".join(f" · {n}" for n in report.notes)
            if report.blocked:
                level = "warn"
                line += f" — BLOQUÉ : {report.blocked['message']}"
            if report.error:
                level = "warn"
                line += f" — {report.error['message']} → {report.error['action']}"
            lines.append(line)
        _flash(conn, cid, "Import terminé.", level, lines)
        return RedirectResponse(f"/c/{cid}", status_code=303)

    @app.get("/c/{cid}/profiles/{pid}", response_class=HTMLResponse)
    def profile_page(request: Request, cid: str, pid: str):
        from paladin.importers import profiles
        from paladin.importers.mapping import MAPPABLE_KEYS, TRANSFORMS

        campaign = _campaign(cid)
        prof = profiles.get(conn, pid)
        source_fields = _profile_source_fields(prof)
        return render(
            request,
            "profile.html",
            campaign=campaign,
            prof=prof,
            keys=MAPPABLE_KEYS,
            transforms=TRANSFORMS,
            source_fields=source_fields,
            flash=_pop_flash(conn, cid),
        )

    def _profile_source_fields(prof: dict) -> list[str]:
        fields = [spec["source"] for spec in prof["mapping"].get("fields", {}).values()]
        fields += prof["mapping"].get("comments_from", [])
        detected = prof["mapping"].get("detected_fields", [])
        return detected + sorted(set(fields) - set(detected))

    @app.post("/c/{cid}/profiles/{pid}/validate")
    async def profile_validate(request: Request, cid: str, pid: str):
        from paladin.importers import profiles
        from paladin.importers.mapping import MAPPABLE_KEYS, MappingError

        form = await _form(request)
        fields = {}
        for key in MAPPABLE_KEYS:
            src = form.get(f"src_{key}", "")
            if src:
                fields[key] = {"source": src, "transform": form.get(f"tr_{key}", "text") or "text"}
        comments = [c for c in form.get("comments_from", "").split("|") if c]
        detected = profiles.get(conn, pid)["mapping"].get("detected_fields", [])
        mapping = {"fields": fields, "constants": {}, "comments_from": comments, "detected_fields": detected}
        try:
            profiles.validate(conn, pid, mapping)
        except MappingError as exc:
            _flash(conn, cid, f"Mapping invalide : {exc}", "warn")
            return RedirectResponse(f"/c/{cid}/profiles/{pid}", status_code=303)
        _flash(conn, cid, "Profil validé. Lancer l'import pour l'appliquer.", "ok")
        return RedirectResponse(f"/c/{cid}", status_code=303)

    @app.post("/c/{cid}/export")
    async def run_export(request: Request, cid: str):
        from paladin.excel.export import export_workbook

        form = await _form(request)
        mode = "final" if form.get("mode") == "final" else "working_copy"
        result = export_workbook(settings, conn, cid, mode=mode)
        if result.status != "verified":
            _flash(
                conn, cid, f"Export non effectué ({result.status}) : {result.error}", "error", [f"→ {result.action}"]
            )
        else:
            s = result.summary
            details = [
                f"{s['cells_written']} cellule(s) écrite(s), {s['findings_up_to_date']} décision(s) à jour dans"
                " l'Excel."
            ]
            if s["blocked"]:
                details += [f"BLOQUÉ {b['sheet']} {b['source_id']} : {b['reason']}" for b in s["blocked"]]
            if s["without_target"]:
                details.append(f"{len(s['without_target'])} finding(s) sans ligne cible.")
            if s["tools_without_sheet"]:
                details.append(f"Outils sans onglet (hors export) : {', '.join(s['tools_without_sheet'])}.")
            _flash(conn, cid, f"Excel à jour : {result.destination}", "ok" if not s["blocked"] else "warn", details)
        return RedirectResponse(f"/c/{cid}", status_code=303)

    @app.post("/c/{cid}/f/{fid}/open-editor")
    async def open_editor(request: Request, cid: str, fid: str):
        from paladin.analysis import _repo_for

        await _form(request)
        command = settings.editor.get("open_command")
        row = conn.execute("SELECT * FROM finding WHERE id = ? AND campaign_id = ?", (fid, cid)).fetchone()
        if not command or row is None or not row["normalized_path"]:
            _flash(conn, cid, "Éditeur non configuré : renseigner [editor] open_command dans paladin.toml.", "warn")
        else:
            repo, rel = _repo_for(conn, row, row["normalized_path"], None)
            path = safe_repo_file(repo, rel) if repo else None
            if path is None:
                _flash(conn, cid, "Fichier introuvable dans les dépôts autorisés.", "warn")
            else:
                args = [
                    a.replace("{path}", str(path)).replace("{line}", str(row["line_number"] or 1))
                    for a in shlex.split(command, posix=True)
                ]
                subprocess.Popen(args)  # noqa: S603 — commande configurée par l'utilisateur, sans shell
        return RedirectResponse(f"/c/{cid}/f/{fid}?{urlencode({'view': 'all'})}", status_code=303)

    return app
