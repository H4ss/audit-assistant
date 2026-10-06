"""Cohérence des fixtures de démo : elles fondent toute la recette."""

import json
import re

from openpyxl import load_workbook

from paladin.demo import fixtures_root, load_demo_descriptor
from paladin.fixtures.workbook import TOOLB_ROWS, build_demo_workbook

ROOT = fixtures_root()


def _issues(version_id: int) -> list[dict]:
    out = []
    for page in sorted((ROOT / "fortify").glob(f"issues_{version_id}_p*.json")):
        out += json.loads(page.read_text(encoding="utf-8"))["data"]
    return out


def test_fortify_release_fixture_has_transport_duplicate_and_total():
    issues = _issues(10042)
    ids = [i["issueInstanceId"] for i in issues]
    assert len(ids) == 14 and len(set(ids)) == 13
    total = json.loads((ROOT / "fortify" / "issues_10042_p1.json").read_text(encoding="utf-8"))["count"]
    assert total == 13


def test_fortify_locations_exist_in_demo_repos():
    descriptor = load_demo_descriptor()
    roots = {r["scanner_roots"][0]: ROOT / r["path"] for r in descriptor["repos"]}
    for issue in _issues(10042):
        full = issue["fullFileName"]
        root = next(r for r in roots if full.startswith(r))
        path = roots[root] / full[len(root):]
        assert path.is_file(), full
        if issue["lineNumber"] is not None:
            lines = path.read_text(encoding="utf-8").splitlines()
            assert 1 <= issue["lineNumber"] <= len(lines), full
            assert lines[issue["lineNumber"] - 1].strip(), f"ligne vide {full}:{issue['lineNumber']}"


def test_fortify_version_selection_cases_present():
    def names(app_id):
        return [v["name"] for v in json.loads((ROOT / "fortify" / f"versions_{app_id}.json").read_text(encoding="utf-8"))["data"]]

    assert names(501).count("release") == 1  # cas nominal
    assert names(502).count("release") == 0  # absente : blocage attendu
    assert names(503).count("release") == 2  # ambiguë : blocage attendu


def test_toolb_excel_and_md_overlap_as_designed():
    md = (ROOT / "toolb" / "toolb_details.md").read_text(encoding="utf-8")
    md_ids = set(re.findall(r"^## (TB-\d{4})", md, flags=re.M))
    xl_ids = {r[0] for r in TOOLB_ROWS}
    assert xl_ids - md_ids == {"TB-0007"}
    assert md_ids - xl_ids == {"TB-0009"}
    assert len(xl_ids & md_ids) == 7


def test_toolc_contains_prompt_injection_and_malformed_row():
    md = (ROOT / "toolc" / "toolc_report.md").read_text(encoding="utf-8")
    assert "ignore all previous instructions" in md
    assert "this row is malformed" in md


def test_demo_workbook_structure(tmp_path):
    path = build_demo_workbook(tmp_path / "wb.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Synthèse", "Fortify", "ToolB"]
    assert wb["Synthèse"]["B3"].value.startswith("=COUNTIF(")
    ws = wb["ToolB"]
    headers = [c.value for c in ws[1]]
    assert "Notes équipe" in headers and "Found in Fortify" in headers
    assert ws["N2"].value == '=IF(J2="","à faire","fait")'
    row_tb8 = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "TB-0008")
    assert row_tb8[9] == "Not an issue"
    assert wb["Fortify"].max_row == 1
