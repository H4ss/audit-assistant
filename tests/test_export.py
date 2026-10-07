"""Export Excel : valeurs exactes, préservation, verrous, conflits, périmé."""

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from tests.conftest import finding_by_source

from paladin.contracts import DISCUSSION_COMMENT, ExportState
from paladin.excel import export as ex
from paladin.review import decisions as d

F1 = "FFFFFFFFFFFFFFFFFFFFFFFF00900001"
F13 = "FFFFFFFFFFFFFFFFFFFFFFFF00900013"


def decide(conn, sid, verdict="TRUE_POSITIVE", comment=None, action="accept"):
    f = finding_by_source(conn, sid)
    return d.record_decision(
        conn, f["id"], expected_revision=f["revision"], action=action, author="me", verdict=verdict, comment=comment
    )


def run(demo, **kw):
    return ex.export_workbook(demo["settings"], demo["conn"], "demo", **kw)


def rows_by_key(ws, key_header):
    headers = [c.value for c in ws[1]]
    k = headers.index(key_header)
    return headers, {r[k]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[k]}


def test_exact_values_metadata_and_preservation(demo):
    conn = demo["conn"]
    decide(conn, F1, "TRUE_POSITIVE", DISCUSSION_COMMENT)
    decide(conn, "FFFFFFFFFFFFFFFFFFFFFFFF00900002", "NOT_AN_ISSUE", DISCUSSION_COMMENT)
    decide(conn, "TB-0001", "TRUE_POSITIVE", "Concaténation SQL confirmée.")
    res = run(demo)
    assert res.status == "verified", res.error
    wb = load_workbook(res.destination)
    assert wb.sheetnames == ["Synthèse", "Fortify", "ToolB"]
    assert wb["Synthèse"]["B3"].value == '=COUNTIF(Fortify!N:N,"True Positive")'  # formule intacte
    assert wb["ToolB"]["N2"].value == '=IF(J2="","à faire","fait")'
    headers, fortify = rows_by_key(wb["Fortify"], "Instance ID")
    assert len(fortify) == 13
    r1 = dict(zip(headers, fortify[F1], strict=True))
    assert r1["analysis result"] == "True Positive" and r1["Analysis result comment"] == DISCUSSION_COMMENT
    assert r1["Category"] == "SQL Injection" and r1["Fortify Category"] == "Input Validation and Representation"
    assert r1["Version name"] == "release" and r1["Line number"] == 12 and r1["CWE"] == "CWE-89"
    assert r1["Found in ToolB"] == "Unknown"  # aucune comparaison lancée : jamais « non détecté »
    r2 = dict(zip(headers, fortify["FFFFFFFFFFFFFFFFFFFFFFFF00900002"], strict=True))
    assert r2["analysis result"] == "Not an issue" and r2["Analysis result comment"] == DISCUSSION_COMMENT
    r13 = dict(zip(headers, fortify[F13], strict=True))
    assert r13["Line number"] is None and r13["analysis result"] is None  # pas de 0 inventé, pas de verdict
    th, toolb = rows_by_key(wb["ToolB"], "Finding ID")
    assert dict(zip(th, toolb["TB-0001"], strict=True))["analysis result"] == "True Positive"
    assert dict(zip(th, toolb["TB-0003"], strict=True))["Notes équipe"] == "voir avec l'équipe front"
    assert dict(zip(th, toolb["TB-0008"], strict=True))["analysis result"] == "Not an issue"  # saisie humaine conservée
    assert all(
        v in (None, "True Positive", "Not an issue")
        for v in [dict(zip(headers, r, strict=True))["analysis result"] for r in fortify.values()]
    )
    assert res.summary["without_target"][0]["source_id"] == "TB-0009"
    assert finding_by_source(conn, F1)["export_state"] == ExportState.EXPORTED
    assert finding_by_source(conn, F13)["export_state"] == ExportState.NOT_EXPORTED


def test_comment_is_written_as_text_not_formula(demo):
    decide(demo["conn"], "TB-0001", "TRUE_POSITIVE", '=HYPERLINK("http://evil","x")')
    res = run(demo)
    ws = load_workbook(res.destination)["ToolB"]
    cell = ws["K2"]
    assert cell.value == '=HYPERLINK("http://evil","x")' and cell.data_type == "s"


def test_human_value_blocks_conflicting_decision_unless_allowed(demo):
    conn = demo["conn"]
    decide(conn, "TB-0008", "TRUE_POSITIVE")
    res = run(demo)
    assert any(b["source_id"] == "TB-0008" for b in res.summary["blocked"])
    assert finding_by_source(conn, "TB-0008")["export_state"] == ExportState.NOT_EXPORTED
    fid = finding_by_source(conn, "TB-0008")["id"]
    res2 = run(demo, allow_overwrite={(fid, "analyst_result"), (fid, "analyst_comment")})
    _, toolb = rows_by_key(load_workbook(res2.destination)["ToolB"], "Finding ID")
    assert toolb["TB-0008"][9] == "True Positive"


def test_second_export_is_idempotent(demo):
    decide(demo["conn"], F1)
    run(demo, mode="final")
    res = run(demo, mode="final")
    assert res.status == "verified" and res.summary["cells_written"] == 0


def test_reordered_workbook_matched_by_key(demo):
    decide(demo["conn"], "TB-0005", "NOT_AN_ISSUE", "Journal applicatif interne.")
    path = ex.target_path(demo["settings"], {"id": "demo", "config": {"target_workbook": "audit_shopapp_demo.xlsx"}})
    wb = load_workbook(path)
    ws = wb["ToolB"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    for i, values in enumerate(reversed(rows), start=2):  # tri externe inversé
        for j, v in enumerate(values, start=1):
            ws.cell(row=i, column=j, value=v)
    wb.save(path)
    res = run(demo, mode="final")
    assert res.status == "verified"
    th, toolb = rows_by_key(load_workbook(path)["ToolB"], "Finding ID")
    assert dict(zip(th, toolb["TB-0005"], strict=True))["analysis result"] == "Not an issue"
    assert dict(zip(th, toolb["TB-0004"], strict=True))["analysis result"] is None


def test_duplicate_key_blocks_only_those_rows(demo):
    decide(demo["conn"], "TB-0001")
    decide(demo["conn"], "TB-0002")
    path = ex.target_path(demo["settings"], {"id": "demo", "config": {"target_workbook": "audit_shopapp_demo.xlsx"}})
    wb = load_workbook(path)
    wb["ToolB"]["A3"] = "TB-0001"  # TB-0002 devient un doublon de TB-0001
    wb.save(path)
    res = run(demo)
    blocked = {b["source_id"] for b in res.summary["blocked"]}
    assert blocked == {"TB-0001"}
    assert any(u["key"] == "TB-0001" for u in res.summary["unmatched_rows"]) is False


def test_locked_workbook_keeps_decisions_pending(demo):
    decide(demo["conn"], F1)
    path = ex.target_path(demo["settings"], {"id": "demo", "config": {"target_workbook": "audit_shopapp_demo.xlsx"}})
    lock = path.with_name("~$" + path.name)
    lock.write_text("x")
    res = run(demo, mode="final")
    assert res.status == "locked" and "Fermer le classeur" in res.action
    assert finding_by_source(demo["conn"], F1)["export_state"] == ExportState.NOT_EXPORTED
    assert finding_by_source(demo["conn"], F1)["current_decision_id"] is not None
    lock.unlink()


def test_replace_permission_error_is_locked(demo, monkeypatch):
    decide(demo["conn"], F1)
    run(demo)  # crée la destination

    def deny(*a, **k):
        raise PermissionError("[WinError 32] fichier utilisé par un autre processus")

    monkeypatch.setattr(ex.Path, "replace", deny)
    decide(demo["conn"], "TB-0001")
    res = run(demo)
    assert res.status == "locked"
    assert finding_by_source(demo["conn"], "TB-0001")["export_state"] == ExportState.NOT_EXPORTED
    assert not list(res.destination.parent.glob(".~paladin-*"))


def test_verification_failure_replaces_nothing(demo, monkeypatch):
    decide(demo["conn"], F1)
    monkeypatch.setattr(ex, "_verify", lambda *a, **k: ["Synthèse!A1 modifiée hors périmètre"])
    res = run(demo)
    assert res.status == "failed" and not res.destination.exists()


def test_workbook_with_chart_is_not_qualified(demo):
    path = ex.target_path(demo["settings"], {"id": "demo", "config": {"target_workbook": "audit_shopapp_demo.xlsx"}})
    wb = load_workbook(path)
    ws = wb["Synthèse"]
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=3, max_row=6))
    ws.add_chart(chart, "D3")
    wb.save(path)
    res = run(demo)
    assert res.status == "failed" and "graphiques" in res.error


def test_undo_after_export_marks_stale_and_reexport_clears(demo):
    conn = demo["conn"]
    decide(conn, "TB-0001", "TRUE_POSITIVE", "ok")
    assert run(demo, mode="final").status == "verified"
    f = finding_by_source(conn, "TB-0001")
    d.undo_last(conn, f["id"], expected_revision=f["revision"], author="me")
    assert finding_by_source(conn, "TB-0001")["export_state"] == ExportState.STALE
    res = run(demo, mode="final")
    _, toolb = rows_by_key(load_workbook(res.destination)["ToolB"], "Finding ID")
    assert toolb["TB-0001"][9] is None and toolb["TB-0001"][10] is None  # valeurs Paladin retirées


def test_investigation_leaves_cells_empty(demo):
    conn = demo["conn"]
    f = finding_by_source(conn, "TB-0002")
    d.record_decision(
        conn,
        f["id"],
        expected_revision=f["revision"],
        action="investigate",
        author="me",
        investigation_question="Paramétrage réel ?",
        investigation_reason="trace partielle",
    )
    res = run(demo)
    _, toolb = rows_by_key(load_workbook(res.destination)["ToolB"], "Finding ID")
    assert toolb["TB-0002"][9] is None


def test_manifest_links_cells_to_decisions(demo):
    import json

    ev = decide(demo["conn"], "TB-0001")
    res = run(demo)
    manifest = json.loads(res.manifest_path.read_text(encoding="utf-8"))
    cell = next(c for c in manifest["cells"] if c["column"] == "analyst_result" and c["new"] == "True Positive")
    assert cell["decision_event_id"] == ev["id"] and cell["sheet"] == "ToolB"
