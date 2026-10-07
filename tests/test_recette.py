"""Recette finale de la tranche (section 24.5), sur une installation neuve et les fixtures.

Parcours : sélection Fortify `release` → import concurrent Excel + MD → onglet proposé pour un
nouvel outil → deux verdicts exacts avec commentaire de discussion → correction annulable → lien
inter-outils confirmé → criticités originales exportées → fermeture / reprise → cellules hors
périmètre préservées.
"""

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from tests.conftest import finding_by_source

from paladin import matching, store
from paladin.config import load_settings
from paladin.contracts import DISCUSSION_COMMENT
from paladin.db import open_database
from paladin.demo import add_simulated_proposals, seed_demo
from paladin.excel import sheets
from paladin.excel.export import export_workbook, target_path
from paladin.importers import profiles
from paladin.importers.pipeline import import_tool
from paladin.review import decisions as d
from paladin.web.app import create_app

F = "FFFFFFFFFFFFFFFFFFFFFFFF00"


def _rows(path, sheet, key):
    ws = load_workbook(path)[sheet]
    headers = [c.value for c in ws[1]]
    k = headers.index(key)
    return {r[k]: dict(zip(headers, r, strict=True)) for r in ws.iter_rows(min_row=2, values_only=True) if r[k]}


def test_slice_acceptance_end_to_end(settings, conn):
    seed_demo(settings, conn)
    # 1. Fortify : version « release » sélectionnée (dev ignorée), identifiant réel conservé
    fortify = import_tool(settings, conn, "demo", "Fortify")
    assert fortify.new == 13 and fortify.completeness == "complete"
    assert finding_by_source(conn, F + "900001")["version_source_id"] == "10042"
    # 2. Concurrent Excel + MD
    toolb = import_tool(settings, conn, "demo", "ToolB")
    assert (toolb.matched, toolb.inventory_only, toolb.details_only) == (7, 1, 1)
    # 3. Nouvel outil sans onglet : profil validé, puis onglet proposé et validé
    toolc = import_tool(settings, conn, "demo", "ToolC")
    profiles.validate(conn, toolc.blocked["profile_id"])
    assert import_tool(settings, conn, "demo", "ToolC").new == 5
    path = target_path(settings, store.get_campaign(conn, "demo"))
    wb = load_workbook(path, read_only=True)
    existing = list(wb.sheetnames)
    wb.close()
    sheets.save_validated(conn, "demo", sheets.propose(conn, "demo", "ToolC", existing), "human")
    add_simulated_proposals(conn)
    # 4. Deux verdicts exacts, chacun avec le commentaire de discussion
    for sid, verdict in ((F + "900001", "TRUE_POSITIVE"), (F + "900002", "NOT_AN_ISSUE")):
        f = finding_by_source(conn, sid)
        d.record_decision(
            conn,
            f["id"],
            expected_revision=f["revision"],
            action="accept",
            author="analyste",
            verdict=verdict,
            comment=DISCUSSION_COMMENT,
        )
    # 5. Correction annulable
    f3 = finding_by_source(conn, F + "900003")
    d.record_decision(
        conn,
        f3["id"],
        expected_revision=f3["revision"],
        action="correct",
        author="analyste",
        verdict="TRUE_POSITIVE",
        correction_category="protection",
    )
    f3 = finding_by_source(conn, F + "900003")
    d.undo_last(conn, f3["id"], expected_revision=f3["revision"], author="analyste")
    assert finding_by_source(conn, F + "900003")["current_decision_id"] is None
    # 6. Lien inter-outils confirmé (aucun verdict propagé)
    matching.compare_all(conn, "demo")
    rel = conn.execute(
        "SELECT r.* FROM finding_relation r JOIN finding a ON a.id = r.finding_a_id JOIN finding b"
        " ON b.id = r.finding_b_id WHERE a.source_id = ? AND b.source_id = ?",
        (F + "900001", "TB-0001"),
    ).fetchone()
    matching.decide(
        conn,
        rel["id"],
        "confirm",
        author="analyste",
        expected_revision=rel["revision"],
        relation_type="same_occurrence",
    )
    assert finding_by_source(conn, "TB-0001")["current_decision_id"] is None
    # 7. Export : valeurs exactes, criticité originale de l'autre outil, nouvel onglet, hors périmètre intact
    res = export_workbook(settings, conn, "demo", mode="final")
    assert res.status == "verified", res.error
    fort = _rows(path, "Fortify", "Instance ID")
    assert fort[F + "900001"]["analysis result"] == "True Positive"
    assert fort[F + "900002"]["analysis result"] == "Not an issue"
    assert {fort[F + s]["Analysis result comment"] for s in ("900001", "900002")} == {DISCUSSION_COMMENT}
    assert fort[F + "900003"]["analysis result"] is None  # correction annulée : rien n'est écrit
    assert (fort[F + "900001"]["Found in ToolB"], fort[F + "900001"]["criticality in ToolB"]) == ("Yes", "High")
    assert _rows(path, "ToolB", "Finding ID")["TB-0001"]["criticality in Fortify"] == "Critical"
    assert len(_rows(path, "ToolC", "Finding ID")) == 5
    wb = load_workbook(path)
    assert wb["Synthèse"]["B3"].value == '=COUNTIF(Fortify!N:N,"True Positive")'
    assert wb["ToolB"]["I4"].value == "voir avec l'équipe front" and wb["ToolB"]["N2"].value.startswith("=IF(")
    # 8. Fermeture / reprise : nouvelle connexion, nouvel espace d'application, tout est retrouvé
    conn.close()
    again = open_database(load_settings(settings.home).db_path)
    try:
        assert finding_by_source(again, F + "900001")["export_state"] == "exported"
        c = TestClient(create_app(load_settings(settings.home)), base_url="http://127.0.0.1:8765")
        page = c.get(f"/c/demo/f/{finding_by_source(again, F + '900002')['id']}?view=all").text
        assert DISCUSSION_COMMENT in page and "Excel : à jour" in page
        assert "Reprendre" in c.get("/c/demo").text
    finally:
        again.close()


def test_wrong_commit_is_flagged_and_excluded_from_batches(demo):
    from paladin import groups

    conn = demo["conn"]
    add_simulated_proposals(conn)
    conn.execute("UPDATE repo SET commit_sha = 'abc999' WHERE name = 'shop-api'")
    f = finding_by_source(conn, F + "900003")
    c = TestClient(create_app(demo["settings"]), base_url="http://127.0.0.1:8765")
    assert "Mauvais commit possible" in c.get(f"/c/demo/f/{f['id']}?view=all").text
    assert groups.commit_mismatch(conn, f).startswith("Commit scanné demo-0001")
    m = groups._member(conn, f, "NOT_AN_ISSUE", None, [])
    assert not m.eligible and "commit" in m.reason


def test_insufficient_evidence_shows_uncertainty_and_next_action(demo):
    conn = demo["conn"]
    add_simulated_proposals(conn)
    f = finding_by_source(conn, F + "900006")
    page = (
        TestClient(create_app(demo["settings"]), base_url="http://127.0.0.1:8765")
        .get(f"/c/demo/f/{f['id']}?view=all")
        .text
    )
    assert "À investiguer / indéterminé" in page and "Manque" in page and "Prochaine action" in page
    assert '<details id="investigate" open' in page  # le panneau « À investiguer » est déjà ouvert
