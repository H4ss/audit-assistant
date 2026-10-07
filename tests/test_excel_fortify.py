"""Fortify sans SSC : l'onglet « Fortify » du classeur cible, au format du modèle Paladin, sert de source."""

import json

from openpyxl import Workbook, load_workbook
from tests.conftest import finding_by_source

from paladin import campaigns
from paladin.excel import manual, reprise
from paladin.excel.export import export_workbook
from paladin.excel.template import headers_for
from paladin.importers import profiles
from paladin.importers.pipeline import import_tool
from paladin.review import decisions as d

ROWS = [  # (Instance ID, catégorie, fichier, ligne, verdict déjà saisi, commentaire déjà saisi)
    ("A" * 32, "SQL Injection", "app/orders.py", 11, "True Positive", "User input concatenated into SQL."),
    ("B" * 32, "SQL Injection", "app/products.py", 13, None, None),
    ("C" * 32, "Command Injection", "app/admin.py", 13, None, None),
]


def _lab(tmp_path):
    (tmp_path / "code" / "app").mkdir(parents=True)
    headers = headers_for("Fortify", ["Fortify"], True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Fortify"
    ws.append(headers)
    for iid, cat, path, line, verdict, comment in ROWS:
        row = dict.fromkeys(headers)
        row.update(
            {
                "Application name": "VulnShop",
                "Version name": "release",
                "Category": cat,
                "Primary location": path.split("/")[-1],
                "Line number": line,
                "Full filename": "/opt/build/vulnshop/" + path,
                "Criticality": "High",
                "Instance ID": iid,
                "analysis result": verdict,
                "Analysis result comment": comment,
            }
        )
        ws.append([row[h] for h in headers])
    wb.create_sheet("Synthèse").append(["TP", '=COUNTIF(Fortify!N:N,"True Positive")'])
    wb.save(tmp_path / "audit.xlsx")
    desc = {
        "id": "vulnshop",
        "name": "VulnShop",
        "target_workbook": "audit.xlsx",
        "repos": [{"name": "vulnshop", "path": "code", "scanner_roots": ["/opt/build/vulnshop/"]}],
        "tools": [
            {
                "label": "Fortify",
                "kind": "excel",
                "sheet_name": "Fortify",
                "sources": [{"role": "findings", "kind": "excel", "path": "@target", "sheet": "Fortify"}],
            }
        ],
    }
    (tmp_path / "campagne.json").write_text(json.dumps(desc), encoding="utf-8")
    return tmp_path / "campagne.json"


def _imported(settings, conn, tmp_path):
    campaigns.create_from_file(settings, conn, _lab(tmp_path))
    blocked = import_tool(settings, conn, "vulnshop", "Fortify").blocked
    mapping = profiles.get(conn, blocked["profile_id"])["mapping"]
    assert mapping["fields"]["source_id"]["source"] == "Instance ID"
    sources = {spec["source"] for spec in mapping["fields"].values()}
    assert not sources & {"analysis result", "Analysis result comment"}  # colonnes de l'analyste : jamais sources
    profiles.validate(conn, blocked["profile_id"])
    assert import_tool(settings, conn, "vulnshop", "Fortify").new == 3


def test_excel_fortify_sheet_round_trip_with_reprise(settings, conn, tmp_path):
    _imported(settings, conn, tmp_path)
    assert finding_by_source(conn, "A" * 32)["normalized_path"].endswith("app/orders.py")
    _rid, plan = reprise.apply(settings, conn, "vulnshop", "me")
    assert [i.key for i in plan.importable] == ["A" * 32]
    f = finding_by_source(conn, "C" * 32)
    d.record_decision(
        conn, f["id"], expected_revision=f["revision"], action="correct", author="me", verdict="TRUE_POSITIVE"
    )
    res = export_workbook(settings, conn, "vulnshop", mode="final")
    assert res.status == "verified", res.error  # clé « Instance ID » tirée du profil validé
    wb = load_workbook(tmp_path / "audit.xlsx")
    rows = {r[10]: r for r in wb["Fortify"].iter_rows(min_row=2, values_only=True)}
    assert rows["A" * 32][13:15] == ("True Positive", "User input concatenated into SQL.")
    assert rows["C" * 32][13] == "True Positive" and rows["B" * 32][13] is None
    assert wb["Synthèse"]["B1"].value == '=COUNTIF(Fortify!N:N,"True Positive")'


def test_calibration_reads_the_analyst_columns_of_a_paladin_workbook(settings, conn, tmp_path):
    _imported(settings, conn, tmp_path)
    plan = manual.scan(conn, "vulnshop", tmp_path / "audit.xlsx")
    assert plan.column_label("comment") == "Analysis result comment"  # pas « Comments » (scanner, vide)
    assert plan.column_label("file") == "Full filename"
    res = manual.apply(conn, "vulnshop", plan, "me")
    ev = d.current_decision(conn, res.imported[0])
    assert ev["comment"] == "User input concatenated into SQL."
    assert export_workbook(settings, conn, "vulnshop", mode="final").status == "verified"
