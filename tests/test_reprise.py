"""Reprise d'un classeur déjà renseigné et changement de classeur cible."""

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from tests.conftest import finding_by_source

from paladin import cli, groups
from paladin.excel import reprise
from paladin.excel.export import export_workbook, target_path
from paladin.review import decisions as d
from paladin.review import memory
from paladin.web.app import create_app


def _path(demo):
    return target_path(demo["settings"], {"id": "demo", "config": {"target_workbook": "audit_shopapp_demo.xlsx"}})


def _fill(demo, values):
    """values : {clé ToolB: (analysis result, comment)} ; ajoute une ligne inconnue TB-9999."""
    path = _path(demo)
    wb = load_workbook(path)
    ws = wb["ToolB"]
    for r in range(2, ws.max_row + 1):
        key = ws.cell(row=r, column=1).value
        if key in values:
            ws.cell(row=r, column=10, value=values[key][0])
            ws.cell(row=r, column=11, value=values[key][1])
    ws.append(["TB-9999", "ShopApp", "X", "x.py", 1, "Low", None, None, None, "True Positive", None])
    wb.save(path)
    return path


def _decide(conn, sid, verdict, comment=None):
    f = finding_by_source(conn, sid)
    d.record_decision(
        conn, f["id"], expected_revision=f["revision"], action="correct", author="me", verdict=verdict, comment=comment
    )


@pytest.fixture()
def filled(demo):
    _decide(demo["conn"], "TB-0004", "NOT_AN_ISSUE")
    _decide(demo["conn"], "TB-0005", "TRUE_POSITIVE", "déjà vu")
    _fill(
        demo,
        {
            "TB-0001": ("TP", "Concaténation SQL"),
            "TB-0002": ("maybe", None),
            "TB-0003": (None, "à voir"),
            "TB-0004": ("True Positive", None),
            "TB-0005": ("True Positive", "déjà vu"),
        },
    )
    return demo


def test_scan_classifies_every_existing_value(filled):
    plan = reprise.scan(filled["settings"], filled["conn"], "demo")
    status = {i.key: i.status for i in plan.items}
    assert status == {
        "TB-0001": "normalisé",
        "TB-0002": "non reconnu",
        "TB-0003": "commentaire seul",
        "TB-0004": "conflit",
        "TB-0005": "identique",
        "TB-0008": "importable",
        "TB-9999": "sans finding",
    }
    assert {i.key for i in plan.importable} == {"TB-0001", "TB-0008"}


def test_apply_creates_traceable_import_decisions_and_export_normalises(filled):
    conn, settings = filled["conn"], filled["settings"]
    rid, _plan = reprise.apply(settings, conn, "demo", "me")
    ev = d.current_decision(conn, finding_by_source(conn, "TB-0008")["id"])
    assert (ev["action"], ev["authority"], ev["batch_id"]) == ("import", "import", rid)
    assert ev["verdict"] == "NOT_AN_ISSUE" and ev["comment"] == "cache key only, no security use"
    assert finding_by_source(conn, "TB-0008")["export_state"] == "exported"  # déjà exact dans l'Excel
    assert finding_by_source(conn, "TB-0001")["export_state"] == "stale"  # « TP » à réécrire
    assert d.current_decision(conn, finding_by_source(conn, "TB-0004")["id"])["verdict"] == "NOT_AN_ISSUE"  # conflit
    assert memory.pilot_metrics(conn, "demo")["decided_import"] == 2
    res = export_workbook(settings, conn, "demo", mode="final")
    assert res.status == "verified"
    rows = {r[0]: r for r in load_workbook(_path(filled))["ToolB"].iter_rows(min_row=2, values_only=True)}
    assert rows["TB-0001"][9] == "True Positive" and rows["TB-0001"][10] == "Concaténation SQL"
    assert rows["TB-0002"][9] == "maybe"  # valeur non reconnue : jamais touchée
    assert any(b["source_id"] == "TB-0004" for b in res.summary["blocked"])  # conflit : revue explicite


def test_undoing_a_reprise_hands_cells_back_to_the_human(filled):
    conn, settings = filled["conn"], filled["settings"]
    rid, _ = reprise.apply(settings, conn, "demo", "me")
    report = groups.undo(conn, rid, "me")
    assert sorted(report["annulés"]) == ["TB-0001", "TB-0008"]
    assert finding_by_source(conn, "TB-0001")["current_decision_id"] is None
    export_workbook(settings, conn, "demo", mode="final")
    rows = {r[0]: r for r in load_workbook(_path(filled))["ToolB"].iter_rows(min_row=2, values_only=True)}
    assert rows["TB-0001"][9] == "TP" and rows["TB-0008"][9] == "Not an issue"  # saisies humaines intactes


def test_import_action_requires_import_authority(demo):
    f = finding_by_source(demo["conn"], "TB-0001")
    with pytest.raises(d.DecisionError):
        d.record_decision(
            demo["conn"],
            f["id"],
            expected_revision=f["revision"],
            action="import",
            author="me",
            verdict="TRUE_POSITIVE",
        )


def test_change_target_warns_about_missing_sheets(demo, tmp_path):
    from openpyxl import Workbook

    other = tmp_path / "autre classeur.xlsx"
    wb = Workbook()
    wb.active.title = "Fortify"
    wb.save(other)
    missing = reprise.change_target(demo["settings"], demo["conn"], "demo", other)
    assert missing == ["ToolB"]
    assert target_path(demo["settings"], {"id": "demo", "config": {"target_workbook": str(other)}}) == other.resolve()
    with pytest.raises(reprise.RepriseError):
        reprise.change_target(demo["settings"], demo["conn"], "demo", tmp_path / "absent.xlsx")


def test_reprise_from_ui_and_cli(filled, capsys, monkeypatch):
    settings = filled["settings"]
    c, t = TestClient(create_app(settings), base_url="http://127.0.0.1:8765"), settings.ui_token()
    page = c.get("/c/demo/reprise").text
    assert "Reprendre 2 décision(s)" in page and "ni True Positive ni Not an issue" in page
    monkeypatch.setenv("PALADIN_HOME", str(settings.home))
    assert cli.main(["excel", "reprise", "--home", str(settings.home)]) == 0
    assert "Aperçu seulement" in capsys.readouterr().out
    r = c.post("/c/demo/reprise", data={"token": t})
    assert "2 décision(s) reprise(s) du classeur" in r.text
