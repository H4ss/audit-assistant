"""Rapprochement inter-outils, projections comparatives et nouvel onglet (sections 21.4 et 22)."""

import json

import pytest
from openpyxl import load_workbook
from tests.conftest import finding_by_source

from paladin import matching, store
from paladin.contracts import FOUND_IN_NO_CONFIRMED, FOUND_IN_PENDING, FOUND_IN_UNKNOWN, FOUND_IN_YES
from paladin.excel import sheets
from paladin.excel.export import export_workbook, target_path
from paladin.importers import profiles
from paladin.importers.pipeline import import_tool
from paladin.review import decisions as d
from paladin.store import ConflictError

F = "FFFFFFFFFFFFFFFFFFFFFFFF00"


@pytest.fixture()
def full(demo):
    """Démo avec les trois outils importés (profil ToolC validé)."""
    conn, settings = demo["conn"], demo["settings"]
    profiles.validate(conn, demo["reports"]["ToolC"].blocked["profile_id"])
    import_tool(settings, conn, "demo", "ToolC")
    return demo


def rel(conn, a_sid, b_sid):
    return conn.execute(
        "SELECT r.* FROM finding_relation r JOIN finding a ON a.id = r.finding_a_id"
        " JOIN finding b ON b.id = r.finding_b_id"
        " WHERE a.source_id = ? AND b.source_id = ?",
        (a_sid, b_sid),
    ).fetchone()


def confirm(conn, a_sid, b_sid, kind="same_occurrence"):
    r = rel(conn, a_sid, b_sid)
    return matching.decide(conn, r["id"], "confirm", author="me", expected_revision=r["revision"], relation_type=kind)


def rows(path, sheet, key):
    ws = load_workbook(path)[sheet]
    headers = [c.value for c in ws[1]]
    k = headers.index(key)
    return {r[k]: dict(zip(headers, r, strict=True)) for r in ws.iter_rows(min_row=2, values_only=True) if r[k]}


# ----------------------------------------------------------------- candidats


def test_candidates_need_the_same_file_and_explain_themselves(full):
    conn = full["conn"]
    res = matching.compare(conn, "demo", "Fortify", "ToolB")
    assert res.comparable and res.new == 11
    exact = rel(conn, F + "900001", "TB-0001")
    assert exact["proposed_type"] == "same_occurrence" and exact["score"] == 100
    assert "même ligne 12" in json.loads(exact["evidence_json"])
    assert any("criticités brutes" in x for x in json.loads(exact["differences_json"]))
    # faux ami : même fichier, même CWE, source constante → jamais proposé « même occurrence »
    assert rel(conn, F + "900012", "TB-0005")["proposed_type"] == "same_root_cause"
    # un-vers-plusieurs
    assert rel(conn, F + "900010", "TB-0005") and rel(conn, F + "900011", "TB-0005")


def test_cwe_or_line_alone_never_makes_a_candidate(full):
    conn = full["conn"]
    a = dict(finding_by_source(conn, F + "900001"))
    b = dict(finding_by_source(conn, "TB-0002"))  # même CWE, autre fichier
    assert matching.score_pair(a, {**b, "line_number": a["line_number"]}) is None


def test_rejection_is_kept_and_change_reproposes(full):
    conn, settings = full["conn"], full["settings"]
    matching.compare(conn, "demo", "Fortify", "ToolB")
    r = rel(conn, F + "900004", "TB-0004")
    matching.decide(conn, r["id"], "reject", author="me", expected_revision=r["revision"])
    again = matching.compare(conn, "demo", "Fortify", "ToolB")
    assert again.new == 0 and rel(conn, F + "900004", "TB-0004")["state"] == "rejected"
    md = settings.campaign_dir("demo") / "inputs" / "toolb" / "toolb_details.md"
    md.write_text(
        md.read_text(encoding="utf-8").replace(
            "- **Severity**: Medium\n- **File**: shop-api/app/search.py",
            "- **Severity**: High\n- **File**: shop-api/app/search.py",
        ),
        encoding="utf-8",
    )
    import_tool(settings, conn, "demo", "ToolB")
    assert matching.compare(conn, "demo", "Fortify", "ToolB").reproposed == 1
    assert rel(conn, F + "900004", "TB-0004")["state"] == "proposed"


def test_confirmed_link_needs_reexam_when_source_changes(full):
    conn, settings = full["conn"], full["settings"]
    matching.compare(conn, "demo", "Fortify", "ToolB")
    confirm(conn, F + "900001", "TB-0001")
    md = settings.campaign_dir("demo") / "inputs" / "toolb" / "toolb_details.md"
    md.write_text(md.read_text(encoding="utf-8").replace("- **Line**: 12", "- **Line**: 13", 1), encoding="utf-8")
    import_tool(settings, conn, "demo", "ToolB")
    assert matching.compare(conn, "demo", "Fortify", "ToolB").reexam >= 1
    assert rel(conn, F + "900001", "TB-0001")["state"] == "reexam_required"


# ----------------------------------------------------------------- décisions


def test_decisions_are_events_with_undo_and_revision(full):
    conn = full["conn"]
    matching.compare(conn, "demo", "Fortify", "ToolB")
    r = rel(conn, F + "900001", "TB-0001")
    confirm(conn, F + "900001", "TB-0001")
    with pytest.raises(ConflictError):
        matching.decide(conn, r["id"], "reject", author="me", expected_revision=r["revision"])
    cur = rel(conn, F + "900001", "TB-0001")
    matching.decide(conn, r["id"], "undo", author="me", expected_revision=cur["revision"])
    assert rel(conn, F + "900001", "TB-0001")["state"] == "proposed"
    actions = [
        e["action"]
        for e in conn.execute("SELECT action FROM relation_event WHERE relation_id = ? ORDER BY created_at", (r["id"],))
    ]
    assert actions == ["propose", "confirm", "undo"]


def test_no_transitivity_and_no_verdict_propagation(full):
    conn = full["conn"]
    matching.compare_all(conn, "demo")
    f1 = finding_by_source(conn, F + "900008")
    d.record_decision(
        conn, f1["id"], expected_revision=f1["revision"], action="accept", author="me", verdict="NOT_AN_ISSUE"
    )
    confirm(conn, F + "900008", "TB-0008")  # Fortify–ToolB
    confirm(conn, "TB-0008", "C-04")  # ToolB–ToolC
    fc = rel(conn, F + "900008", "C-04")  # Fortify–ToolC : reste proposé, jamais déduit
    assert fc["state"] == "proposed"
    assert finding_by_source(conn, "TB-0008")["current_decision_id"] is None  # aucun verdict propagé
    assert finding_by_source(conn, "C-04")["current_decision_id"] is None


def test_batch_confirms_a_frozen_list_with_one_event_each(full):
    conn = full["conn"]
    matching.compare(conn, "demo", "Fortify", "ToolB")
    exact = matching.exact_candidates(conn, "demo")
    assert len(exact) == 6 and all(e["a_sid"] and e["b_sid"] for e in exact)
    batch = matching.confirm_batch(conn, "demo", [e["id"] for e in exact], "me")
    n = conn.execute("SELECT COUNT(*) FROM relation_event WHERE batch_id = ?", (batch,)).fetchone()[0]
    assert n == 6
    with pytest.raises(matching.MatchingError):
        matching.confirm_batch(conn, "demo", [exact[0]["id"]], "me")  # liste périmée : déjà confirmé


# ----------------------------------------------------------------- projections


def test_projection_policy(full):
    conn = full["conn"]
    fortify = store.get_tool(conn, "demo", "Fortify")
    toolb = store.get_tool(conn, "demo", "ToolB")
    toolc = store.get_tool(conn, "demo", "ToolC")
    f1 = finding_by_source(conn, F + "900001")
    assert matching.projection(conn, f1, toolb["id"]).found_in == FOUND_IN_UNKNOWN  # pas encore comparé
    matching.compare_all(conn, "demo")
    assert matching.projection(conn, f1, toolb["id"]).found_in == FOUND_IN_PENDING
    confirm(conn, F + "900001", "TB-0001")
    p = matching.projection(conn, f1, toolb["id"])
    assert (p.found_in, p.criticality) == (FOUND_IN_YES, "High")
    confirm(conn, F + "900010", "TB-0005")
    confirm(conn, F + "900011", "TB-0005")
    tb5 = finding_by_source(conn, "TB-0005")
    multi = matching.projection(conn, tb5, fortify["id"])
    assert multi.criticality == f"{F}900010: Medium; {F}900011: Medium"  # toutes, avec leurs IDs, sans maximum
    r = rel(conn, F + "900012", "TB-0005")
    matching.decide(
        conn, r["id"], "confirm", author="me", expected_revision=r["revision"], relation_type="same_root_cause"
    )
    f12 = finding_by_source(conn, F + "900012")
    assert matching.projection(conn, f12, toolb["id"]).found_in == FOUND_IN_NO_CONFIRMED  # cause commune ≠ trouvé
    # corpus ToolC partiel : sans lien confirmé, jamais « non détecté »
    f13 = finding_by_source(conn, F + "900013")
    assert matching.projection(conn, f13, toolc["id"]).found_in == FOUND_IN_UNKNOWN


# ----------------------------------------------------------------- export


def test_export_writes_comparative_cells_and_goes_stale_on_undo(full):
    conn, settings = full["conn"], full["settings"]
    matching.compare(conn, "demo", "Fortify", "ToolB")
    f1 = finding_by_source(conn, F + "900001")
    d.record_decision(
        conn, f1["id"], expected_revision=f1["revision"], action="accept", author="me", verdict="TRUE_POSITIVE"
    )
    confirm(conn, F + "900001", "TB-0001")
    confirm(conn, F + "900010", "TB-0005")
    confirm(conn, F + "900011", "TB-0005")
    res = export_workbook(settings, conn, "demo", mode="final")
    assert res.status == "verified", res.error
    path = target_path(settings, store.get_campaign(conn, "demo"))
    fort = rows(path, "Fortify", "Instance ID")
    assert (fort[F + "900001"]["Found in ToolB"], fort[F + "900001"]["criticality in ToolB"]) == ("Yes", "High")
    assert fort[F + "900002"]["Found in ToolB"] == "Pending review"
    tb = rows(path, "ToolB", "Finding ID")
    assert tb["TB-0005"]["criticality in Fortify"] == f"{F}900010: Medium; {F}900011: Medium"
    assert tb["TB-0007"]["Found in Fortify"] == "No confirmed match"
    manifest = json.loads(res.manifest_path.read_text(encoding="utf-8"))
    assert any(c["column"] == "found_in:ToolB" for c in manifest["cells"])
    r = rel(conn, F + "900001", "TB-0001")
    matching.decide(conn, r["id"], "undo", author="me", expected_revision=r["revision"])
    assert finding_by_source(conn, F + "900001")["export_state"] == "stale"
    export_workbook(settings, conn, "demo", mode="final")
    assert rows(path, "Fortify", "Instance ID")[F + "900001"]["Found in ToolB"] == "Pending review"


def test_human_value_in_comparative_cell_is_preserved(full):
    conn, settings = full["conn"], full["settings"]
    path = target_path(settings, store.get_campaign(conn, "demo"))
    wb = load_workbook(path)
    wb["ToolB"]["L2"] = "Oui (vu en réunion)"
    wb.save(path)
    matching.compare(conn, "demo", "Fortify", "ToolB")
    res = export_workbook(settings, conn, "demo", mode="final")
    assert rows(path, "ToolB", "Finding ID")["TB-0001"]["Found in Fortify"] == "Oui (vu en réunion)"
    assert any(p["source_id"] == "TB-0001" for p in res.summary["preserved_human_values"])


# ----------------------------------------------------------------- nouvel onglet


def test_new_tool_sheet_proposal_validation_and_idempotent_creation(full):
    conn, settings = full["conn"], full["settings"]
    path = target_path(settings, store.get_campaign(conn, "demo"))
    wb_ro = load_workbook(path, read_only=True)
    existing = list(wb_ro.sheetnames)
    wb_ro.close()  # Windows : un fichier ouvert ne peut pas être remplacé
    prop = sheets.propose(conn, "demo", "ToolC", existing)
    keys = [c.key for c in prop.columns]
    assert keys[0] == "source_id" and "fortify_category" not in keys and "primary_rule_id" not in keys
    assert {"analyst_result", "analyst_comment", "found_in_fortify", "criticality_in_toolb"} <= set(keys)
    edited = sheets.apply_edits(
        prop,
        {
            "keep_application_name": "1",
            "keep_version_name": "1",
            "keep_category": "1",
            "header_category": "Contrôle",
            "keep_full_filename": "1",
            "keep_criticality_raw": "1",
            "sheet_name": "ToolC",
        },
    )
    assert "line_number" not in [c.key for c in edited.columns]  # retirée (non cochée)
    sheets.validate(edited, existing)
    sheets.save_validated(conn, "demo", edited, "human")
    with pytest.raises(sheets.SheetSchemaError):
        sheets.validate(edited, [*existing, "ToolC"])  # collision détectée
    f = finding_by_source(conn, "C-01")
    d.record_decision(
        conn,
        f["id"],
        expected_revision=f["revision"],
        action="accept",
        author="me",
        verdict="NOT_AN_ISSUE",
        comment="security appetite to be discussed",
    )
    for _ in range(2):  # deux exports : pas de doublon de lignes
        res = export_workbook(settings, conn, "demo", mode="final")
        assert res.status == "verified", res.error
    wb = load_workbook(path)
    assert wb.sheetnames[-1] == "ToolC" and wb.sheetnames[:3] == ["Synthèse", "Fortify", "ToolB"]
    toolc = rows(path, "ToolC", "ID ToolC")
    assert len(toolc) == 5 and wb["ToolC"].max_row == 6
    assert toolc["C-01"]["Contrôle"] == "hardcoded-credential"
    assert toolc["C-01"]["analysis result"] == "Not an issue"
    assert toolc["C-01"]["Analysis result comment"] == "security appetite to be discussed"
    assert toolc["C-01"]["Found in Fortify"] == "Unknown"


def test_schema_evolution_never_removes_columns(full):
    conn = full["conn"]
    prop = sheets.propose(conn, "demo", "ToolC", [])
    smaller = prop.model_copy(update={"columns": [c for c in prop.columns if c.key != "category"]})
    with pytest.raises(sheets.SheetSchemaError):
        sheets.validate(smaller, [], previous=prop)


def test_missing_comparative_columns_added_on_request(full):
    conn, settings = full["conn"], full["settings"]
    cfg = store.get_campaign(conn, "demo")["config"]
    cfg["sheet_extensions"] = {"Fortify": ["Found in ToolC", "criticality in ToolC"]}
    store.update_campaign_config(conn, "demo", cfg)
    matching.compare_all(conn, "demo")
    res = export_workbook(settings, conn, "demo", mode="final")
    assert res.status == "verified", res.error
    ws = load_workbook(target_path(settings, store.get_campaign(conn, "demo")))["Fortify"]
    headers = [c.value for c in ws[1]]
    assert headers[-2:] == ["Found in ToolC", "criticality in ToolC"]
    assert headers[:2] == ["Application name", "Version name"]  # rien n'est déplacé
